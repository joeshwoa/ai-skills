"""Editable-Word renderer: normalised Doc → native .docx (python-docx).

Everything stays editable: real Heading 1–3 styles (navigation pane + TOC work), a TOC field Word
fills in on open, true tables with repeating header rows, page-number fields, native bullets, and for
Arabic a right-to-left document (section, paragraphs, runs and tables flagged RTL, Arabic
complex-script fonts). Images and charts are pictures; everything else is text.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Emu

from pdk.text import runs as md_runs


def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


class DocxRenderer:
    def __init__(self, doc):
        self.d = doc
        self.tk = doc.tk
        self.rtl = doc.rtl
        self.fig = 0
        self.tab = 0

    # ================================================================== setup
    def render(self, out: Path):
        self.doc = Document()
        self._page_setup(self.doc.sections[0])
        self._styles()
        self._settings()
        lay = self.d.layout
        if lay == "flow":
            self._cover()
            if self.d.toc and self.d.outline:
                self._toc()
            body = self.doc
            self._headers_footers(self.doc.sections[-1])
            self.blocks(self.d.blocks, body)
        else:
            getattr(self, f"layout_{lay}")()
        if self.d.credits():
            p = self.para(self.doc, ("حقوق الصور" if self.rtl else "Image credits"), size=8, bold=True, color=self.tk["muted"], space_before=18)
            for c in self.d.credits():
                self.para(self.doc, c, size=7.5, color=self.tk["muted"], space_after=1)
        cp = self.doc.core_properties
        cp.title = self.d.meta.get("title", "")
        cp.author = self.d.meta.get("author", "")
        cp.subject = self.d.meta.get("subtitle", "")
        self.doc.save(str(out))
        return str(out)

    def _page_setup(self, sec):
        w, h = self.d.page_mm
        sec.page_width, sec.page_height = Mm(w), Mm(h)
        sec.orientation = WD_ORIENT.LANDSCAPE if self.d.orientation == "landscape" else WD_ORIENT.PORTRAIT
        t, r, b, l = self.d.margins_mm
        sec.top_margin, sec.right_margin, sec.bottom_margin, sec.left_margin = Mm(t), Mm(r), Mm(b), Mm(l)
        sec.header_distance = Mm(9)
        sec.footer_distance = Mm(9)
        if self.rtl:
            sec._sectPr.append(_el("w:bidi"))

    def _fonts(self, rpr_owner, font):
        rpr = rpr_owner.get_or_add_rPr() if hasattr(rpr_owner, "get_or_add_rPr") else rpr_owner
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = _el("w:rFonts")
            rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), font)

    def _styles(self):
        st = self.doc.styles
        tk = self.tk
        normal = st["Normal"]
        normal.font.name = tk["body_font"]
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = rgb(tk["ink"])
        self._fonts(normal.element, tk["body_font"])
        rpr = normal.element.get_or_add_rPr()
        sz = rpr.find(qn("w:szCs")) or _el("w:szCs", **{"w:val": "21"})
        rpr.append(sz)
        if self.rtl:
            rpr.append(_el("w:rtl"))
            ppr = normal.element.get_or_add_pPr()
            ppr.append(_el("w:bidi"))
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.45 if self.rtl else 1.3
        for lvl, size, color in ((1, 20, tk["ink"]), (2, 14.5, tk["accent"]), (3, 12, tk["ink_soft"])):
            hs = st[f"Heading {lvl}"]
            hs.font.name = tk["head_font"]
            hs.font.size = Pt(size)
            hs.font.bold = True
            hs.font.color.rgb = rgb(color)
            self._fonts(hs.element, tk["head_font"])
            hr = hs.element.get_or_add_rPr()
            hr.append(_el("w:szCs", **{"w:val": str(int(size * 2))}))
            hr.append(_el("w:bCs"))
            hs.paragraph_format.space_before = Pt(18 if lvl == 1 else 14 if lvl == 2 else 10)
            hs.paragraph_format.space_after = Pt(6 if lvl < 3 else 3)
            hs.paragraph_format.keep_with_next = True
            if self.rtl:
                hs.element.get_or_add_pPr().append(_el("w:bidi"))
        for name in ("List Bullet", "List Number"):
            s = st[name]
            s.font.name = tk["body_font"]
            if self.rtl:
                s.element.get_or_add_pPr().append(_el("w:bidi"))
        cap = st["Caption"]
        cap.font.size = Pt(8.5)
        cap.font.italic = False
        cap.font.color.rgb = rgb(tk["muted"])
        cap.font.name = tk["body_font"]

    def _settings(self):
        s = self.doc.settings.element
        if self.d.toc:
            s.append(_el("w:updateFields", **{"w:val": "true"}))    # Word offers to fill the TOC on open
        lang = s.find(qn("w:themeFontLang"))
        if lang is None:
            lang = _el("w:themeFontLang")
            s.append(lang)
        lang.set(qn("w:val"), "ar-EG" if self.rtl else "en-US")
        lang.set(qn("w:bidi"), "ar-EG")

    # ================================================================== primitives
    def para(self, container, text="", *, style=None, size=None, bold=False, color=None, align=None,
             space_after=None, space_before=None, font=None, italic=False, md=True, keep=False):
        p = container.add_paragraph(style=style) if style else container.add_paragraph()
        if text not in (None, ""):
            self.add_runs(p, str(text), size=size, bold=bold, color=color, font=font, italic=italic, md=md)
        self._bidi(p)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "justify":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif align == "end":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if self.rtl else WD_ALIGN_PARAGRAPH.RIGHT
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        if space_before is not None:
            p.paragraph_format.space_before = Pt(space_before)
        if keep:
            p.paragraph_format.keep_with_next = True
        return p

    def _bidi(self, p):
        if self.rtl:
            ppr = p._p.get_or_add_pPr()
            if ppr.find(qn("w:bidi")) is None:
                ppr.insert(0, _el("w:bidi"))

    def add_runs(self, p, text, *, size=None, bold=False, color=None, font=None, italic=False, md=True):
        parts = md_runs(text) if md else [(text, {})]
        out = []
        for t, stl in parts:
            r = p.add_run(t)
            if size:
                r.font.size = Pt(size)
                r._r.get_or_add_rPr().append(_el("w:szCs", **{"w:val": str(int(size * 2))}))
            if bold or stl.get("bold"):
                r.font.bold = True
                r._r.get_or_add_rPr().append(_el("w:bCs"))
            if italic:
                r.font.italic = True
            col = self.tk["accent"] if stl.get("accent") else color
            if col:
                r.font.color.rgb = rgb(col)
            f = "Consolas" if stl.get("code") else font
            if f:
                r.font.name = f
                self._fonts(r._r, f)
            if self.rtl and not stl.get("code"):
                r._r.get_or_add_rPr().append(_el("w:rtl"))
            out.append(r)
        return out

    def shade(self, cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:shd")):
            tcPr.remove(old)
        tcPr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))

    def cell_borders(self, cell, **sides):
        """sides: top/bottom/left/right/start/end = (size_eighths, color) or None for none."""
        tcPr = cell._tc.get_or_add_tcPr()
        b = tcPr.find(qn("w:tcBorders"))
        if b is None:
            b = _el("w:tcBorders")
            tcPr.append(b)
        for side, val in sides.items():
            if side == "start":
                side = "right" if self.rtl else "left"
            if side == "end":
                side = "left" if self.rtl else "right"
            e = _el(f"w:{side}", **({"w:val": "nil"} if val is None else {"w:val": "single", "w:sz": str(val[0]), "w:color": val[1], "w:space": "0"}))
            b.append(e)

    def cell_margins(self, cell, mm=2.5):
        tcPr = cell._tc.get_or_add_tcPr()
        m = _el("w:tcMar")
        for s in ("top", "bottom", "left", "right"):
            m.append(_el(f"w:{s}", **{"w:w": str(int(mm * 56.7)), "w:type": "dxa"}))
        tcPr.append(m)

    def table(self, container, rows, cols, *, widths_pct=None, borders=False, rtl=None):
        t = container.add_table(rows=rows, cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        if (self.rtl if rtl is None else rtl):
            tblPr.append(_el("w:bidiVisual"))
        if not borders:
            tb = _el("w:tblBorders")
            for s in ("top", "left", "bottom", "right", "insideH", "insideV"):
                tb.append(_el(f"w:{s}", **{"w:val": "nil"}))
            tblPr.append(tb)
        tw = self.text_width_mm(container)
        widths_pct = widths_pct or [100 / cols] * cols
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = _el("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:w"), str(int(tw * 56.7)))
        tblW.set(qn("w:type"), "dxa")
        for j, col in enumerate(t.columns):
            col.width = Mm(tw * widths_pct[j] / 100)          # gridCol (LibreOffice/Google Docs read this)
            for c in col.cells:
                c.width = Mm(tw * widths_pct[j] / 100)
        return t

    def text_width_mm(self, container=None):
        if container is not None and hasattr(container, "_tc"):
            try:
                return container.width / 36000
            except Exception:
                pass
        w = self.d.page_mm[0] - self.d.margins_mm[1] - self.d.margins_mm[3]
        return w

    def cell_para(self, cell, text="", **kw):
        p = cell.paragraphs[0] if (len(cell.paragraphs) == 1 and not cell.paragraphs[0].text and not cell.paragraphs[0].runs) else cell.add_paragraph()
        if text not in (None, ""):
            self.add_runs(p, str(text), size=kw.get("size"), bold=kw.get("bold", False), color=kw.get("color"),
                          font=kw.get("font"), md=kw.get("md", True))
        self._bidi(p)
        if kw.get("align") == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kw.get("align") == "end":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if self.rtl else WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(kw.get("space_after", 2))
        return p

    def picture(self, container, path, width_mm, *, align="center"):
        if not path:
            return None
        if hasattr(container, "_tc"):
            p = self.cell_para(container)
        else:
            p = container.add_paragraph()
            self._bidi(p)
        p.add_run().add_picture(str(path), width=Mm(width_mm))
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        return p

    def crop(self, path, aspect):
        """Centre-crop an image to width/height = aspect (covers, section banners)."""
        if not path:
            return path
        from PIL import Image
        im = Image.open(path)
        w, h = im.size
        if abs(w / h - aspect) < 0.02:
            return path
        if w / h > aspect:
            nw = int(h * aspect)
            box = ((w - nw) // 2, 0, (w - nw) // 2 + nw, h)
        else:
            nh = int(w / aspect)
            top = int((h - nh) * 0.35)
            box = (0, top, w, top + nh)
        out = Path(path).with_name(Path(path).stem + f"-crop{int(aspect * 100)}.jpg")
        if not out.exists():
            im.crop(box).convert("RGB").save(out, quality=88)
        return str(out)

    def field(self, p, instr, placeholder=""):
        r = p.add_run()
        r._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
        r2 = p.add_run()
        it = _el("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f" {instr} "
        r2._r.append(it)
        r3 = p.add_run()
        r3._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
        r4 = p.add_run(placeholder)
        r5 = p.add_run()
        r5._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
        return r4

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ================================================================== cover, toc, header/footer
    def _cover(self):
        d, c, m, tk = self.d, self.d.cover, self.d.meta, self.tk
        style = c.get("style", "minimal")
        if style == "none":
            return
        if style == "masthead":
            self.para(self.doc, m.get("org", ""), size=8.5, bold=True, color=tk["accent"], space_after=0)
            self.para(self.doc, m.get("title", ""), size=28, bold=True, font=tk["head_font"], space_after=2, md=False)
            p = self.para(self.doc, " · ".join(v for v in (m.get("issue", ""), m.get("date", "")) if v), size=9, color=tk["muted"], space_after=12)
            bdr = _el("w:pBdr")
            bdr.append(_el("w:bottom", **{"w:val": "single", "w:sz": "12", "w:space": "4", "w:color": tk["ink"]}))
            p._p.get_or_add_pPr().append(bdr)
            return
        tw = self.text_width_mm()
        logo = d.asset(c.get("logo") or ({"path": d.logo} if d.logo else None)) if (c.get("logo") or d.logo) else None
        if logo:
            self.picture(self.doc, logo, 34, align=None)
        if style in ("photo", "art", "band", "split"):
            ref = c.get("image") or {"art": c.get("art", "blobs" if style == "art" else "geo"), "dark": style != "band", "seed": 3}
            img = self.crop(d.asset(ref, size=(2000, 1300)), 16 / 10.5)   # keep the cover on ONE page
            self.picture(self.doc, img, tw)
        else:
            for _ in range(5):
                self.para(self.doc, "", space_after=12)
            bar = self.table(self.doc, 1, 1, widths_pct=[100])
            cell = bar.cell(0, 0)
            cell.width = Mm(28)
            self.shade(cell, tk["accent"])
            self.cell_para(cell, " ", size=2)
        kick = c.get("kicker") or m.get("kicker")
        if kick:
            self.para(self.doc, kick.upper() if not self.rtl else kick, size=10, bold=True, color=tk["accent"], space_before=18, space_after=4)
        self.para(self.doc, m.get("title", ""), size=30, bold=True, color=tk["ink"], font=tk["head_font"],
                  space_before=6 if kick else 22, space_after=8, md=False)
        if m.get("subtitle"):
            self.para(self.doc, m["subtitle"], size=13, color=tk["ink_soft"], space_after=18)
        meta_rows = [(self.d.label(lab), m[k]) for k, lab in (("prepared_for", "prepared_for"), ("author", "prepared_by"),
                                                             ("date", "date"), ("version", "version")) if m.get(k)]
        if meta_rows:
            t = self.table(self.doc, 1, len(meta_rows))
            for j, (lab, val) in enumerate(meta_rows):
                cell = t.cell(0, j)
                self.cell_borders(cell, top=(12, tk["accent"]))
                self.cell_para(cell, lab.upper() if not self.rtl else lab, size=7.5, bold=True, color=tk["muted"])
                self.cell_para(cell, val, size=10)
        if m.get("confidential"):
            self.para(self.doc, self.d.label("confidential"), size=9, bold=True, color="B42318", space_before=16)
        self.doc.add_section(WD_SECTION.NEW_PAGE)

    def _toc(self):
        tk = self.tk
        self.para(self.doc, self.d.label("contents"), size=20, bold=True, font=tk["head_font"], space_after=12)
        outline = self.d.outline
        tw = self.text_width_mm()
        for i, o in enumerate(outline):
            p = self.doc.add_paragraph()
            self._bidi(p)
            pf = p.paragraph_format
            pf.space_after = Pt(3)
            pf.tab_stops.add_tab_stop(Mm(tw), alignment=0 if self.rtl else 2, leader=1)   # right tab, dots
            if o["level"] > 1:
                pf.left_indent = Mm(6) if not self.rtl else None
                if self.rtl:
                    pf.right_indent = Mm(6)
            if i == 0:
                r = p.add_run()
                r._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
                r2 = p.add_run()
                it = _el("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = ' TOC \\o "1-2" \\h \\z \\u '
                r2._r.append(it)
                r3 = p.add_run()
                r3._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
            txt = (o["num"] + "  " if o["num"] else "") + o["text"]
            self.add_runs(p, txt, size=11 if o["level"] == 1 else 10, bold=o["level"] == 1, md=False,
                          color=tk["ink"] if o["level"] == 1 else tk["ink_soft"])
            p.add_run("\t")
            if i == len(outline) - 1:
                r = p.add_run()
                r._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
        note = "← حدّث الفهرس (F9) لأرقام الصفحات" if self.rtl else "Update the table (F9 / right-click → Update field) for page numbers"
        self.para(self.doc, note, size=7.5, color=tk["muted"], space_before=6)
        self.doc.add_section(WD_SECTION.NEW_PAGE)

    def _headers_footers(self, sec):
        d, tk = self.d, self.tk
        sec.different_first_page_header_footer = False
        sec.header.is_linked_to_previous = False      # keep the cover/TOC sections clean
        sec.footer.is_linked_to_previous = False
        hdr_text = d.spec.get("header") if d.spec.get("header") is not None else (d.meta.get("title", "") if d.preset.get("header") else "")
        if hdr_text:
            hp = sec.header.paragraphs[0]
            self.add_runs(hp, hdr_text, size=8, color=tk["muted"], md=False)
            self._bidi(hp)
            pPr = hp._p.get_or_add_pPr()
            bdr = _el("w:pBdr")
            bdr.append(_el("w:bottom", **{"w:val": "single", "w:sz": "4", "w:space": "4", "w:color": tk["line"]}))
            pPr.append(bdr)
        fp = sec.footer.paragraphs[0]
        self._bidi(fp)
        foot = d.spec.get("footer", d.meta.get("org", ""))
        tw = self.text_width_mm()
        fp.paragraph_format.tab_stops.add_tab_stop(Mm(tw), alignment=0 if self.rtl else 2)
        if foot:
            self.add_runs(fp, foot, size=8, color=tk["muted"], md=False)
        fp.add_run("\t")
        r = fp.add_run(f"{d.L['page']} ")
        r.font.size = Pt(8)
        r.font.color.rgb = rgb(tk["muted"])
        ph = self.field(fp, "PAGE", "1")
        ph.font.size = Pt(8)
        r = fp.add_run(f" {d.L['of']} ")
        r.font.size = Pt(8)
        r.font.color.rgb = rgb(tk["muted"])
        ph = self.field(fp, "NUMPAGES", "1")
        ph.font.size = Pt(8)

    # ================================================================== blocks
    def blocks(self, blocks, container):
        for b in blocks or []:
            fn = getattr(self, f"b_{b['type']}", None)
            if fn:
                fn(b, container)

    def b_heading(self, b, c):
        lvl = min(int(b.get("_level", b.get("level", 1))), 3)
        txt = (b["_num"] + "  " if b.get("_num") else "") + b.get("text", "")
        if hasattr(c, "_tc"):
            self.cell_para(c, txt, bold=True, size={1: 14, 2: 12, 3: 11}[lvl], color=self.tk["accent"])
            return
        p = self.doc.add_paragraph(style=f"Heading {lvl}")
        self.add_runs(p, txt, md=False)
        self._bidi(p)

    def b_section(self, b, c):
        if len(self.doc.paragraphs) > 3:
            self.page_break()
        sn = b.get("number") or b.get("_num") or ""
        if sn:
            self.para(self.doc, str(sn), size=44, bold=True, color=self.tk["accent_soft"] if False else self.tk["accent"], font=self.tk["head_font"], space_after=0)
        p = self.doc.add_paragraph(style="Heading 1")
        self.add_runs(p, ((b["_num"] + "  ") if b.get("_num") and not sn else "") + b.get("title", b.get("text", "")), md=False)
        self._bidi(p)
        if b.get("intro"):
            self.para(self.doc, b["intro"], size=12, color=self.tk["ink_soft"], space_after=10)
        if b.get("image"):
            self.picture(self.doc, self.crop(self.d.asset(b["image"], size=(1800, 800)), 2.6), self.text_width_mm())

    def _p(self, c, text, **kw):
        if hasattr(c, "_tc"):
            return self.cell_para(c, text, **kw)
        return self.para(c, text, **kw)

    def b_paragraph(self, b, c):
        lead = b.get("style") == "lead"
        self._p(c, b.get("text", ""), size=12 if lead else None, color=self.tk["ink_soft"] if lead else None)

    def b_bullets(self, b, c, ordered=False):
        style = "List Number" if ordered or b.get("style") == "number" else "List Bullet"
        for it in b.get("items", []):
            if hasattr(c, "_tc"):
                p = self.cell_para(c, "")
                p.style = self.doc.styles[style]
            else:
                p = c.add_paragraph(style=style)
                self._bidi(p)
            if isinstance(it, dict):
                self.add_runs(p, it.get("point", it.get("title", "")), bold=True)
                if it.get("detail"):
                    self.add_runs(p, " — " + it["detail"], color=self.tk["ink_soft"])
            else:
                self.add_runs(p, str(it))
            p.paragraph_format.space_after = Pt(3)

    def b_numbered(self, b, c):
        self.b_bullets(b, c, True)

    def b_checklist(self, b, c):
        for it in b.get("items", []):
            done = isinstance(it, dict) and it.get("done")
            txt = it.get("text", "") if isinstance(it, dict) else it
            p = self._p(c, "")
            r = p.add_run("☑  " if done else "☐  ")
            r.font.color.rgb = rgb(self.tk["accent"])
            self.add_runs(p, txt)
            p.paragraph_format.space_after = Pt(3)

    def b_callout(self, b, c):
        kind = b.get("kind", "info")
        fill = {"warning": "FFF4E5", "success": "E9F7EF", "note": self.tk["paper_alt"]}.get(kind, self.tk["accent_soft"])
        col = {"warning": "B45309", "success": "1E7A46", "note": self.tk["ink_soft"]}.get(kind, self.tk["accent"])
        t = self.table(c, 1, 1)
        cell = t.cell(0, 0)
        self.shade(cell, fill)
        self.cell_borders(cell, start=(24, col))
        self.cell_margins(cell, 3)
        self.cell_para(cell, (b.get("title") or self.d.label(kind)).upper() if not self.rtl else (b.get("title") or self.d.label(kind)),
                       size=8.5, bold=True, color=col)
        if b.get("blocks"):
            self.blocks(b["blocks"], cell)
        else:
            self.cell_para(cell, b.get("text", ""))
        self._gap(c)

    def _gap(self, c, pt=4):
        if not hasattr(c, "_tc"):
            p = c.add_paragraph()
            p.paragraph_format.space_after = Pt(pt)
            p.paragraph_format.line_spacing = Pt(pt)

    def b_quote(self, b, c):
        t = self.table(c, 1, 1)
        cell = t.cell(0, 0)
        self.cell_borders(cell, start=(24, self.tk["accent"]))
        self.cell_margins(cell, 3)
        p = self.cell_para(cell, b.get("text", ""), size=12, color=self.tk["ink_soft"])
        for r in p.runs:
            r.font.italic = True
        if b.get("author"):
            self.cell_para(cell, "— " + b["author"], size=9, bold=True, color=self.tk["muted"])
        self._gap(c)

    def b_pull_quote(self, b, c):
        t = self.table(c, 1, 1)
        cell = t.cell(0, 0)
        self.cell_borders(cell, top=(12, self.tk["accent"]), bottom=(12, self.tk["accent"]))
        self.cell_para(cell, b.get("text", ""), size=16, bold=True, color=self.tk["accent"])
        self._gap(c)

    def b_table(self, b, c):
        heads = b.get("headers", [])
        rows = b.get("rows", [])
        n = len(heads) or (len(rows[0]) if rows else 1)
        t = self.table(c, len(rows) + (1 if heads else 0), n, widths_pct=b.get("widths"))
        tk = self.tk
        minimal = b.get("style") == "minimal"
        r0 = 0
        if heads:
            trPr = t.rows[0]._tr.get_or_add_trPr()
            trPr.append(_el("w:tblHeader"))
            for j, h in enumerate(heads):
                cell = t.cell(0, j)
                if minimal:
                    self.cell_borders(cell, bottom=(12, tk["ink"]))
                else:
                    self.shade(cell, tk["accent"])
                self.cell_para(cell, h, size=9, bold=True, color=tk["ink"] if minimal else tk["on_accent"])
            r0 = 1
        for i, row in enumerate(rows):
            last = b.get("total_row") and i == len(rows) - 1
            for j, v in enumerate(row):
                cell = t.cell(i + r0, j)
                if i % 2 == 1 and not minimal and not last:
                    self.shade(cell, tk["paper_alt"])
                self.cell_borders(cell, bottom=(4, tk["line"]) if not last else None, top=(12, tk["ink"]) if last else None)
                num = bool(re.fullmatch(r"[\s$€£%+\-−.,0-9٠-٩]+", str(v))) and re.search(r"\d", str(v))
                self.cell_para(cell, v, size=9.5, bold=bool(last), align="end" if num else None)
        self.tab += 1
        if b.get("caption"):
            self._p(c, f"**{self.d.label('table')} {self.tab}.** {b['caption']}", size=8.5, color=tk["muted"])
        else:
            self._gap(c)

    def b_image(self, b, c):
        size = b.get("width", "full")
        p = self.d.asset(b.get("image"), size=(1800, 1000) if size == "full" else (1000, 800))
        if not p:
            return
        tw = self.text_width_mm(c)
        w = {"full": tw, "half": tw / 2, "third": tw / 3}.get(size, tw)
        self.picture(c, p, w)
        self.fig += 1
        if b.get("caption"):
            self._p(c, f"**{self.d.label('figure')} {self.fig}.** {b['caption']}", size=8.5, color=self.tk["muted"], align="center")

    def b_chart(self, b, c):
        spec = dict(b["chart"]) if isinstance(b.get("chart"), dict) else {k: v for k, v in b.items() if k not in ("type",)}
        spec.setdefault("type", "bar")
        p = self.d.asset({"chart": spec, "width_in": 7.0, "height_in": float(b.get("height_in", 3.4))})
        self.picture(c, p, self.text_width_mm(c))
        self.fig += 1
        if b.get("caption"):
            self._p(c, f"**{self.d.label('figure')} {self.fig}.** {b['caption']}", size=8.5, color=self.tk["muted"], align="center")

    def b_stats(self, b, c):
        items = b.get("items", [])
        t = self.table(c, 1, max(len(items), 1))
        for j, s in enumerate(items):
            cell = t.cell(0, j)
            self.cell_borders(cell, top=(24, self.tk["accent"]))
            self.cell_para(cell, s.get("value", ""), size=20, bold=True, color=self.tk["accent"], font=self.tk["head_font"])
            self.cell_para(cell, s.get("label", ""), size=8.5, color=self.tk["ink_soft"])
        self._gap(c, 8)

    def b_timeline(self, b, c):
        items = b.get("items", [])
        t = self.table(c, len(items), 2, widths_pct=[22, 78])
        for i, it in enumerate(items):
            a, z = t.cell(i, 0), t.cell(i, 1)
            self.cell_para(a, it.get("date", ""), size=9, bold=True, color=self.tk["accent"])
            self.cell_borders(z, start=(12, self.tk["accent"]))
            self.cell_para(z, it.get("title", ""), bold=True)
            if it.get("text"):
                self.cell_para(z, it["text"], size=9.5, color=self.tk["ink_soft"])
        self._gap(c)

    def b_steps(self, b, c):
        items = b.get("items", [])
        t = self.table(c, len(items), 2, widths_pct=[9, 91])
        for i, it in enumerate(items):
            a, z = t.cell(i, 0), t.cell(i, 1)
            self.shade(a, self.tk["accent"])
            a.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            self.cell_para(a, str(i + 1), size=11, bold=True, color=self.tk["on_accent"], align="center")
            self.cell_margins(z, 2)
            self.cell_para(z, it.get("title", ""), bold=True)
            if it.get("text"):
                self.cell_para(z, it["text"], size=9.5, color=self.tk["ink_soft"])
        self._gap(c)

    def b_two_column(self, b, c):
        ratio = {"1:2": [34, 66], "2:1": [66, 34]}.get(b.get("ratio", ""), [50, 50])
        t = self.table(c, 1, 2, widths_pct=ratio)
        for j, key in enumerate(("left", "right")):
            cell = t.cell(0, j)
            self.cell_margins(cell, 2.5)
            self.blocks(b.get(key), cell)
        self._gap(c)

    def b_box(self, b, c):
        t = self.table(c, 1, 1)
        cell = t.cell(0, 0)
        for s in ("top", "bottom", "left", "right"):
            self.cell_borders(cell, **{s: (8, self.tk["line"])})
        self.cell_margins(cell, 3.5)
        if b.get("title"):
            self.cell_para(cell, b["title"], bold=True, color=self.tk["accent"], size=11)
        self.blocks(b.get("blocks"), cell)
        self._gap(c)

    def b_cards(self, b, c):
        items = b.get("items", [])
        cols = b.get("cols") or (2 if len(items) in (2, 4) else 3)
        rows = (len(items) + cols - 1) // cols
        t = self.table(c, rows, cols)
        for i, it in enumerate(items):
            cell = t.cell(i // cols, i % cols)
            self.shade(cell, self.tk["paper_alt"])
            self.cell_margins(cell, 3)
            self.cell_borders(cell, top=(24, self.d.tk["paper"]), bottom=(24, self.d.tk["paper"]),
                              left=(24, self.d.tk["paper"]), right=(24, self.d.tk["paper"]))
            if it.get("icon"):
                ic = self.d.asset({"icon": it["icon"]} if isinstance(it["icon"], str) else it["icon"], color=self.tk["accent"])
                if ic:
                    self.picture(cell, ic, 8, align=None)
            self.cell_para(cell, it.get("title", ""), bold=True, size=10.5)
            self.cell_para(cell, it.get("text", ""), size=9, color=self.tk["ink_soft"])
        self._gap(c)

    def b_definition(self, b, c):
        t = self.table(c, 1, 1)
        cell = t.cell(0, 0)
        for s in ("top", "bottom", "left", "right"):
            self.cell_borders(cell, **{s: (6, self.tk["line"])})
        self.cell_margins(cell, 3)
        self.cell_para(cell, self.d.label("key_term"), size=7.5, bold=True, color=self.tk["muted"])
        self.cell_para(cell, b.get("term", ""), size=12, bold=True, color=self.tk["accent"])
        self.cell_para(cell, b.get("text", ""))
        if b.get("example"):
            self.cell_para(cell, f"*{self.d.label('example')}:* {b['example']}", size=9.5)
        self._gap(c)

    def b_code(self, b, c):
        t = self.table(c, 1, 1, rtl=False)
        cell = t.cell(0, 0)
        self.shade(cell, "1E2433")
        self.cell_margins(cell, 3)
        colors = {"Comment": "8B9AAB", "Keyword": "7EE0FF", "Name.Builtin": "7EE0FF", "String": "E8C547",
                  "Number": "C3A6FF", "Name.Function": "9BE36D", "Name.Class": "9BE36D", "Operator": "FF6B9D"}
        try:
            from pygments import lex
            from pygments.lexers import get_lexer_by_name, guess_lexer
            lx = get_lexer_by_name(b["lang"]) if b.get("lang") else guess_lexer(b.get("code", ""))
            toks = list(lex(b.get("code", "").rstrip("\n"), lx))
        except Exception:
            toks = [(None, b.get("code", ""))]
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        if toks and toks[-1][1].endswith("\n"):
            toks[-1] = (toks[-1][0], toks[-1][1].rstrip("\n"))
        for tt, val in toks:
            parts = val.split("\n")
            for k, part in enumerate(parts):
                if k > 0:
                    p = cell.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.15
                if not part:
                    continue
                r = p.add_run(part)
                r.font.name = "Consolas"
                self._fonts(r._r, "Consolas")
                r.font.size = Pt(9)
                col = "E6EDF3"
                if tt is not None:
                    s = str(tt).replace("Token.", "")
                    for key, cc in colors.items():
                        if s.startswith(key):
                            col = cc
                            break
                r.font.color.rgb = rgb(col)
        if cell.paragraphs and not cell.paragraphs[-1].text:
            pass
        self._gap(c)

    def b_kv(self, b, c):
        items = b.get("items", [])
        t = self.table(c, len(items), 2, widths_pct=[30, 70])
        for i, (k, v) in enumerate(items):
            self.cell_para(t.cell(i, 0), k, bold=True, size=9, color=self.tk["muted"])
            self.cell_para(t.cell(i, 1), v)
        self._gap(c)

    def b_contact(self, b, c):
        self.b_kv(b, c)

    def b_signature(self, b, c):
        parties = b.get("parties", [])
        self._gap(c, 18)
        t = self.table(c, 1, max(len(parties), 1))
        for j, p in enumerate(parties):
            cell = t.cell(0, j)
            self.cell_para(cell, " ", space_after=26)
            pp = self.cell_para(cell, p.get("name", ""), bold=True)
            pp.paragraph_format.space_before = Pt(2)
            self.cell_borders(cell, bottom=None)
            pPr = cell.paragraphs[0]._p.get_or_add_pPr()
            bdr = _el("w:pBdr")
            bdr.append(_el("w:bottom", **{"w:val": "single", "w:sz": "8", "w:space": "1", "w:color": self.tk["ink"]}))
            pPr.append(bdr)
            role = p.get("role", "")
            if p.get("date_line", True):
                role += f"   {self.d.label('date')}: ________"
            self.cell_para(cell, role, size=8.5, color=self.tk["muted"])

    def b_page_break(self, b, c):
        if not hasattr(c, "_tc"):
            self.page_break()

    def b_divider(self, b, c):
        p = self._p(c, "")
        pPr = p._p.get_or_add_pPr()
        bdr = _el("w:pBdr")
        bdr.append(_el("w:bottom", **{"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": self.tk["line"]}))
        pPr.append(bdr)

    def b_spacer(self, b, c):
        p = self._p(c, "")
        p.paragraph_format.space_after = Pt(float(b.get("mm", 8)) * 2.83)

    def b_toc(self, b, c):
        pass

    def b_gallery(self, b, c):
        items = b.get("items", [])
        cols = 3
        rows = (len(items) + cols - 1) // cols
        t = self.table(c, rows, cols)
        cw = self.text_width_mm(c) / cols - 3
        for i, it in enumerate(items):
            cell = t.cell(i // cols, i % cols)
            self.picture(cell, self.d.asset(it.get("image"), size=(1000, 750)), cw)
            if it.get("caption"):
                self.cell_para(cell, it["caption"], size=8, color=self.tk["muted"])
        self._gap(c)

    def b_note_lines(self, b, c):
        """Writing lines. A table (one row per line) — adjacent bordered paragraphs would merge into one box."""
        n = int(b.get("lines", 5))
        t = self.table(c, n, 1)
        for i in range(n):
            cell = t.cell(i, 0)
            self.cell_borders(cell, bottom=(4, self.tk["muted"]))
            tr = t.rows[i]._tr.get_or_add_trPr()
            tr.append(_el("w:trHeight", **{"w:val": str(int(9 * 56.7)), "w:hRule": "atLeast"}))
            self.cell_para(cell, "", space_after=0)

    def b_schedule(self, b, c):
        items = b.get("items", [])
        t = self.table(c, len(items), 3, widths_pct=[16, 12, 72])
        for i, r in enumerate(items):
            self.cell_para(t.cell(i, 0), r.get("time", ""), bold=True, color=self.tk["accent"], size=9.5)
            self.cell_para(t.cell(i, 1), r.get("duration", ""), color=self.tk["muted"], size=9)
            z = t.cell(i, 2)
            self.cell_para(z, r.get("title", ""), bold=True)
            if r.get("text"):
                self.cell_para(z, r["text"], size=9.5, color=self.tk["ink_soft"])
            for cell in (t.cell(i, 0), t.cell(i, 1), z):
                self.cell_borders(cell, bottom=(4, self.tk["line"]))
        self._gap(c)

    def b_rubric(self, b, c):
        levels = b.get("levels", [])
        crit = b.get("criteria", [])
        t = self.table(c, len(crit) + 1, len(levels) + 1, widths_pct=[24] + [76 / max(len(levels), 1)] * len(levels))
        self.shade(t.cell(0, 0), self.tk["accent"])
        self.cell_para(t.cell(0, 0), self.d.label("criteria"), bold=True, size=9, color=self.tk["on_accent"])
        for j, l in enumerate(levels):
            self.shade(t.cell(0, j + 1), self.tk["accent"])
            self.cell_para(t.cell(0, j + 1), l, bold=True, size=9, color=self.tk["on_accent"])
        for i, cr in enumerate(crit):
            self.cell_para(t.cell(i + 1, 0), cr.get("name", ""), bold=True, size=9.5)
            if cr.get("weight"):
                self.cell_para(t.cell(i + 1, 0), str(cr["weight"]), size=8, color=self.tk["muted"])
            for j, x in enumerate(cr.get("descriptors", [])[:len(levels)]):
                self.cell_para(t.cell(i + 1, j + 1), x, size=9)
            for j in range(len(levels) + 1):
                self.cell_borders(t.cell(i + 1, j), bottom=(4, self.tk["line"]))
        self._gap(c)

    def b_items(self, b, c):
        T = self.d.invoice_totals(b)
        cur = b.get("currency", self.d.spec.get("currency", ""))
        L = self.d.label
        self.b_table({"headers": [L("item"), L("qty"), L("price"), L("amount")],
                      "rows": [[f"**{r[0]}**" + (f" — {r[1]}" if r[1] else ""), f"{r[2]:g}", self.d.money(r[3], cur), self.d.money(r[4], cur)] for r in T["rows"]],
                      "widths": [52, 10, 19, 19]}, c)
        lines = [(L("subtotal"), self.d.money(T["subtotal"], cur))]
        if T["discount"]:
            lines.append((L("discount"), "− " + self.d.money(T["discount"], cur)))
        if T["tax_pct"]:
            lines.append((f"{L('tax')} ({T['tax_pct']:g}%)", self.d.money(T["tax"], cur)))
        lines.append((L("total"), self.d.money(T["total"], cur)))
        t = self.table(c, len(lines), 3, widths_pct=[52, 26, 22])
        for i, (k, v) in enumerate(lines):
            grand = i == len(lines) - 1
            a, z = t.cell(i, 1), t.cell(i, 2)
            if grand:
                self.shade(a, self.tk["accent"])
                self.shade(z, self.tk["accent"])
            col = self.tk["on_accent"] if grand else None
            self.cell_para(a, k, bold=grand, color=col, size=11 if grand else 9.5)
            self.cell_para(z, v, bold=grand, color=col, size=11 if grand else 9.5, align="end")
        self._gap(c, 8)

    def b_questions(self, b, c):
        show = b.get("show_answers", self.d.type == "answer_key")
        letters = "أبجدهو" if self.rtl else "ABCDEF"
        for k, q in enumerate(b.get("items", []), start=int(b.get("start", 1))):
            kind = q.get("kind", "mcq" if q.get("options") else "short")
            p = self._p(c, "")
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
            self.add_runs(p, f"{k}. ", bold=True, color=self.tk["accent"])
            self.add_runs(p, q.get("q", ""), bold=True)
            if q.get("points"):
                self.add_runs(p, f"   ({q['points']} {self.d.label('points')})", size=8.5, color=self.tk["muted"], md=False)
            if kind in ("mcq", "tf"):
                opts = q.get("options") or [self.d.label("true"), self.d.label("false")]
                ans = q.get("answer")
                for j, o in enumerate(opts):
                    right = show and (ans == j or ans == o)
                    pp = self._p(c, "")
                    pp.paragraph_format.space_after = Pt(1)
                    if self.rtl:
                        pp.paragraph_format.right_indent = Mm(7)
                    else:
                        pp.paragraph_format.left_indent = Mm(7)
                    self.add_runs(pp, ("●" if right else "○") + f" {letters[j]})  ", color=self.tk["accent"] if right else self.tk["ink_soft"], md=False)
                    self.add_runs(pp, str(o), bold=right)
            elif show and q.get("answer") is not None:
                pp = self._p(c, f"✓ {q['answer']}", color="1E7A46", bold=True)
            elif kind != "fill":
                n = int(q.get("lines", 2 if kind == "short" else 6))
                self.b_note_lines({"lines": n}, c)
            if show and q.get("explanation"):
                self._p(c, q["explanation"], size=9, color=self.tk["ink_soft"])

    def b_experience(self, b, c):
        for x in b.get("items", []):
            dates = " – ".join(v for v in (x.get("from", ""), x.get("to", "")) if v)
            t = self.table(c, 1, 2, widths_pct=[72, 28])
            self.cell_para(t.cell(0, 0), x.get("role", x.get("degree", "")), bold=True, size=10.5)
            self.cell_para(t.cell(0, 1), dates, size=8.5, color=self.tk["muted"], align="end")
            org = " · ".join(v for v in (x.get("org", x.get("school", "")), x.get("place", "")) if v)
            if org:
                self._p(c, org, size=9, bold=True, color=self.tk["accent"], space_after=2)
            for i in x.get("bullets", []):
                p = c.add_paragraph(style="List Bullet") if not hasattr(c, "_tc") else self.cell_para(c, "")
                self._bidi(p)
                self.add_runs(p, i, size=9.5)
                p.paragraph_format.space_after = Pt(1)
            if x.get("text"):
                self._p(c, x["text"], size=9.5)

    def b_education(self, b, c):
        self.b_experience(b, c)

    def b_skills(self, b, c):
        for g in b.get("groups", []):
            p = self._p(c, "")
            if g.get("group"):
                self.add_runs(p, g["group"] + ": ", bold=True, size=9.5)
            self.add_runs(p, " · ".join(g.get("items", [])), size=9.5)

    # ================================================================== layouts
    def _letterhead(self):
        d, m = self.d, self.d.meta
        t = self.table(self.doc, 1, 2, widths_pct=[50, 50])
        a, z = t.cell(0, 0), t.cell(0, 1)
        logo = d.asset({"path": d.logo}) if d.logo else (d.asset(m["logo"]) if m.get("logo") else None)
        if logo:
            self.picture(a, logo, 32, align=None)
        else:
            self.cell_para(a, m.get("org", ""), size=15, bold=True, color=self.tk["accent"])
        if m.get("org_address"):
            for line in str(m["org_address"]).split("\n"):
                self.cell_para(z, line, size=8.5, color=self.tk["muted"], align="end", space_after=0)
        for cell in (a, z):
            self.cell_borders(cell, bottom=(16, self.tk["accent"]))
        self._gap(self.doc, 14)

    def layout_letter(self):
        d, m = self.d, self.d.meta
        self._letterhead()
        self.para(self.doc, m.get("date", ""), color=self.tk["muted"], space_after=12)
        if m.get("to"):
            for line in str(m["to"]).split("\n"):
                self.para(self.doc, line, space_after=0)
            self.para(self.doc, "", space_after=6)
        if m.get("subject"):
            self.para(self.doc, f"{d.label('subject')}: {m['subject']}", bold=True, space_after=10)
        self.blocks(d.blocks, self.doc)

    def layout_memo(self):
        d, m = self.d, self.d.meta
        if d.logo or m.get("org"):
            self._letterhead()
        self.para(self.doc, m.get("title") or ("مذكرة" if self.rtl else "MEMO"), size=24, bold=True, color=self.tk["accent"], space_after=8)
        rows = [(d.label(k), m[k]) for k in ("to", "from", "cc", "date", "subject") if m.get(k)]
        if d.type == "minutes" and m.get("attendees"):
            rows.append((d.label("attendees"), ", ".join(m["attendees"]) if isinstance(m["attendees"], list) else m["attendees"]))
        self.b_kv({"items": rows}, self.doc)
        self.b_divider({}, self.doc)
        self._headers_footers(self.doc.sections[0]) if d.preset.get("header") else None
        self.blocks(d.blocks, self.doc)

    def layout_invoice(self):
        d, m = self.d, self.d.meta
        t = self.table(self.doc, 1, 2, widths_pct=[50, 50])
        a, z = t.cell(0, 0), t.cell(0, 1)
        if d.logo:
            self.picture(a, d.asset({"path": d.logo}), 34, align=None)
        else:
            self.cell_para(a, m.get("org", ""), size=16, bold=True, color=self.tk["accent"])
        kind = d.label(d.type if d.type in ("invoice", "quotation", "receipt") else "invoice")
        self.cell_para(z, kind.upper() if not self.rtl else kind, size=26, bold=True, color=self.tk["accent"], align="end")
        if m.get("number"):
            self.cell_para(z, f"{d.label('number')} {m['number']}", size=9.5, color=self.tk["muted"], align="end")
        self._gap(self.doc, 12)
        t = self.table(self.doc, 1, 3)
        cols = [(d.label("from"), m.get("from", m.get("org", ""))), (d.label("bill_to"), m.get("bill_to", "")),
                (d.label("date"), m.get("date", "") + (f"\n{d.label('due')}: {m['due']}" if m.get("due") else ""))]
        for j, (k, v) in enumerate(cols):
            cell = t.cell(0, j)
            self.cell_para(cell, k.upper() if not self.rtl else k, size=7.5, bold=True, color=self.tk["muted"])
            for line in str(v).split("\n"):
                self.cell_para(cell, line, size=9.5, space_after=0)
        self._gap(self.doc, 12)
        self.blocks(d.blocks, self.doc)

    def layout_cv(self):
        d, m = self.d, self.d.meta
        sec = self.doc.sections[0]
        sec.top_margin = sec.bottom_margin = Mm(12)
        sec.left_margin = sec.right_margin = Mm(12)
        side_blocks = [b for b in d.blocks if b.get("column") == "side" or b["type"] in ("contact", "skills")]
        main_blocks = [b for b in d.blocks if b not in side_blocks]
        t = self.table(self.doc, 1, 2, widths_pct=[33, 67])
        side, main = t.cell(0, 0), t.cell(0, 1)
        self.shade(side, self.tk["stage_bg"])
        self.cell_margins(side, 5)
        self.cell_margins(main, 5)
        on = self.tk["stage_text"]
        if m.get("photo"):
            self.picture(side, d.asset(m["photo"]), 36)
        for b in side_blocks:
            self.cell_para(side, (b.get("title") or d.label(b["type"])).upper() if not self.rtl else (b.get("title") or d.label(b["type"])),
                           size=8, bold=True, color=self.tk["accent_2"] if self.tk["style"] == "dark" else "FFFFFF")
            if b["type"] == "contact":
                for k, v in b.get("items", []):
                    self.cell_para(side, k, size=7.5, color=self.tk["muted"] if self.tk["style"] == "light" else on, space_after=0)
                    self.cell_para(side, v, size=9, color=on)
            elif b["type"] == "skills":
                for g in b.get("groups", []):
                    if g.get("group"):
                        self.cell_para(side, g["group"], size=8.5, bold=True, color=on, space_after=1)
                    self.cell_para(side, " · ".join(g.get("items", [])), size=9, color=on)
            else:
                for line in [b.get("text", "")] + [str(i) for i in b.get("items", [])]:
                    if line:
                        self.cell_para(side, line, size=9, color=on)
            self.cell_para(side, " ", size=4)
        self.cell_para(main, m.get("title") or m.get("name", ""), size=24, bold=True, font=self.tk["head_font"], space_after=0)
        self.cell_para(main, m.get("subtitle") or m.get("role", ""), size=12, bold=True, color=self.tk["accent"], space_after=8)
        for b in main_blocks:
            if b["type"] in ("experience", "education") or b.get("title"):
                p = self.cell_para(main, (b.get("title") or d.label(b["type"])).upper() if not self.rtl else (b.get("title") or d.label(b["type"])),
                                   size=10, bold=True, color=self.tk["accent"])
                p.paragraph_format.space_before = Pt(10)
                pPr = p._p.get_or_add_pPr()
                bdr = _el("w:pBdr")
                bdr.append(_el("w:bottom", **{"w:val": "single", "w:sz": "4", "w:space": "1", "w:color": self.tk["line"]}))
                pPr.append(bdr)
            if b["type"] in ("experience", "education"):
                for x in b.get("items", []):
                    dates = " – ".join(v for v in (x.get("from", ""), x.get("to", "")) if v)
                    p = self.cell_para(main, x.get("role", x.get("degree", "")), bold=True, size=10.5, space_after=0)
                    if dates:
                        self.add_runs(p, f"   {dates}", size=8.5, color=self.tk["muted"], md=False)
                    org = " · ".join(v for v in (x.get("org", x.get("school", "")), x.get("place", "")) if v)
                    if org:
                        self.cell_para(main, org, size=9, bold=True, color=self.tk["accent"], space_after=1)
                    for i in x.get("bullets", []):
                        self.cell_para(main, "• " + i, size=9.2, space_after=1)
                    if x.get("text"):
                        self.cell_para(main, x["text"], size=9.2)
            else:
                self.blocks([b], main)

    def layout_certificate(self):
        d, m = self.d, self.d.meta
        sec = self.doc.sections[0]
        pg = _el("w:pgBorders", **{"w:offsetFrom": "page"})
        for s in ("top", "left", "bottom", "right"):
            pg.append(_el(f"w:{s}", **{"w:val": "double", "w:sz": "18", "w:space": "24", "w:color": self.tk["accent"]}))
        sec._sectPr.append(pg)
        for _ in range(2):
            self.para(self.doc, "", space_after=10)
        if d.logo:
            self.picture(self.doc, d.asset({"path": d.logo}), 30)
        self.para(self.doc, m.get("org", ""), size=11, bold=True, color=self.tk["accent_2"], align="center", space_before=10)
        self.para(self.doc, m.get("title") or d.label("certificate_of"), size=32, bold=True, color=self.tk["accent"], align="center",
                  font=self.tk["head_font"], space_after=14)
        self.para(self.doc, m.get("presented_to") or d.label("presented_to"), size=11, color=self.tk["muted"], align="center")
        self.para(self.doc, m.get("recipient", ""), size=28, bold=True, align="center", font=self.tk["head_font"], space_after=8)
        self.para(self.doc, m.get("reason", ""), size=12, color=self.tk["ink_soft"], align="center", space_after=6)
        self.para(self.doc, m.get("date", ""), size=10, color=self.tk["muted"], align="center", space_after=30)
        sigs = m.get("signatures", [])
        if sigs:
            t = self.table(self.doc, 1, len(sigs))
            for j, s in enumerate(sigs):
                cell = t.cell(0, j)
                self.cell_borders(cell, top=(8, self.tk["ink"]))
                self.cell_para(cell, s.get("name", ""), bold=True, align="center")
                self.cell_para(cell, s.get("role", ""), size=9, color=self.tk["muted"], align="center")

    def _band(self, kicker, title):
        d = self.d
        t = self.table(self.doc, 1, 2, widths_pct=[78, 22])
        a, z = t.cell(0, 0), t.cell(0, 1)
        if kicker:
            self.cell_para(a, kicker, size=8.5, bold=True, color=self.tk["accent"], space_after=0)
        self.cell_para(a, title, size=19, bold=True, font=self.tk["head_font"])
        if d.logo:
            self.picture(z, d.asset({"path": d.logo}), 24)
        for cell in (a, z):
            self.cell_borders(cell, bottom=(20, self.tk["accent"]))
        self._gap(self.doc, 8)

    def _fields(self):
        L = self.d.label
        t = self.table(self.doc, 1, 3, widths_pct=[50, 25, 25])
        for j, k in enumerate(("name", "class", "date")):
            cell = t.cell(0, j)
            self.cell_borders(cell, bottom=(8, self.tk["ink_soft"]))
            self.cell_para(cell, f"{L(k)}:", size=9, color=self.tk["muted"])
        self._gap(self.doc, 8)

    def layout_handout(self):
        d, m = self.d, self.d.meta
        self._headers_footers(self.doc.sections[0])
        self._band(m.get("kicker") or m.get("course", ""), m.get("title", ""))
        if d.type in ("worksheet", "lab_sheet") or m.get("student_fields"):
            self._fields()
        if m.get("instructions"):
            self.b_callout({"kind": "note", "title": "التعليمات" if self.rtl else "Instructions", "text": m["instructions"]}, self.doc)
        self.blocks(d.blocks, self.doc)

    def layout_exam(self):
        d, m = self.d, self.d.meta
        self._headers_footers(self.doc.sections[0])
        total = sum(float(q.get("points", 0)) for b in d.blocks if b["type"] == "questions" for q in b.get("items", []))
        info = " · ".join(v for v in (m.get("course", ""), (f"{d.label('duration')}: {m['duration']}" if m.get("duration") else ""),
                                      (f"{d.label('score')}: ____ / {total:g}" if total else "")) if v)
        self._band(info, m.get("title", ""))
        if d.type != "answer_key":
            self._fields()
        if m.get("instructions"):
            self.b_callout({"kind": "note", "title": "التعليمات" if self.rtl else "Instructions", "text": m["instructions"]}, self.doc)
        self.blocks(d.blocks, self.doc)

    def layout_flyer(self):
        d, m, c = self.d, self.d.meta, self.d.cover
        img = d.asset(c.get("image") or {"art": "mesh", "dark": True}, size=(1800, 1100))
        self.picture(self.doc, img, self.text_width_mm())
        self.para(self.doc, m.get("title", ""), size=28, bold=True, font=self.tk["head_font"], space_before=8, space_after=4)
        if m.get("subtitle"):
            self.para(self.doc, m["subtitle"], size=13, color=self.tk["ink_soft"], space_after=10)
        self.blocks(d.blocks, self.doc)
        cta = m.get("cta")
        if cta:
            t = self.table(self.doc, 1, 2, widths_pct=[62, 38])
            for cell in t.rows[0].cells:
                self.shade(cell, self.tk["accent"])
                self.cell_margins(cell, 4)
            self.cell_para(t.cell(0, 0), cta.get("text", ""), size=15, bold=True, color=self.tk["on_accent"])
            if cta.get("detail"):
                self.cell_para(t.cell(0, 0), cta["detail"], size=9.5, color=self.tk["on_accent"])
            self.cell_para(t.cell(0, 1), cta.get("contact", ""), size=12, bold=True, color=self.tk["on_accent"], align="end")

    def layout_one_pager(self):
        self.layout_flyer()
