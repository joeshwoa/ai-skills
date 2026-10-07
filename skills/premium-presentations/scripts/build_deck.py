#!/usr/bin/env python3
"""
build_deck.py — turns a JSON "deck spec" into a fully editable, themed .pptx.

Usage:
    python3 build_deck.py deck_spec.json output.pptx

Why python-pptx directly (not HTML->PPTX conversion): HTML/LibreOffice
conversion flattens most text into images or uneditable boxes. Building
the deck natively with python-pptx keeps every title, bullet, and table
cell as real, editable PowerPoint text — which is the #1 requirement
for this skill ("fully editable"). The HTML renderer (render_preview.py)
reads the *same* deck spec + theme so the web preview and the .pptx
always match, and gives you rich CSS animations for previewing/PDF.

ALWAYS run scripts/validate_deck.py on the output before handing a deck
over — see SKILL.md's "Check it before you hand it over" step. It caught
real bugs during this skill's own development (a too-tight section-number
box, an empty-footer crash) that eyeballing a few slides missed.

Deck spec schema: see ../references/deck_spec_schema.md
"""
import json
import math
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

sys.path.insert(0, str(Path(__file__).parent))
from themes import get_theme, render_background_png

SLIDE_W = Inches(13.333)  # 16:9
SLIDE_H = Inches(7.5)

# One background PNG per (theme, variant) is plenty to render every slide
# in a deck sharply — generating one per slide would just waste build time
# for a pixel-identical result.
_BG_CACHE = {}


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def set_background(slide, theme, variant):
    """Stamps the theme's generated gradient+grid+glow art as a full-bleed
    picture behind the slide's content, instead of a flat solid fill —
    flat fills are what make a deck look like a default template."""
    key = (theme["id"], variant)
    if key not in _BG_CACHE:
        tmp_dir = Path(tempfile.gettempdir()) / "premium_presentations_bg"
        tmp_dir.mkdir(parents=True, exist_ok=True)
        path = tmp_dir / f"{theme['id']}_{variant}.png"
        if not path.exists():
            render_background_png(theme, variant, str(path))
        _BG_CACHE[key] = str(path)
    slide.shapes.add_picture(_BG_CACHE[key], 0, 0, width=SLIDE_W, height=SLIDE_H)


def add_transition(slide, kind="fade", duration_ms=700):
    """python-pptx has no public API for slide transitions, so we inject the
    <p:transition> element directly into the slide XML. 'fade' and 'push'
    render correctly in PowerPoint, Keynote, and Google Slides."""
    sld = slide._element
    existing = sld.find(qn('p:transition'))
    if existing is not None:
        sld.remove(existing)
    transition = etree.SubElement(sld, qn('p:transition'))
    transition.set('spd', 'med')
    transition.set('dur', str(duration_ms))
    effect_tag = {
        "fade": qn('p:fade'),
        "push": qn('p:push'),
        "wipe": qn('p:wipe'),
        "cut": None,
    }.get(kind)
    if effect_tag is not None:
        etree.SubElement(transition, effect_tag)


def add_textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=None):
    """Every textbox in this file goes through here. This matters more than
    it looks: if word_wrap / auto_size are left at python-pptx's defaults,
    LibreOffice (and some other renderers) silently auto-shrinks the box to
    fit its text and re-centers it around the *original* box's center —
    which quietly breaks left-alignment even though nothing in the XML looks
    wrong. Pinning word_wrap=True and auto_size=NONE makes every renderer
    respect the exact box you placed."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    if anchor is not None:
        tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return box, tf


def set_text(tf, text, align=None):
    """Use instead of `tf.text = text`. python-pptx's tf.text setter leaves a
    paragraph with ZERO runs when text is an empty string — style_run() then
    crashes on paragraphs[0].runs[0]. This always leaves exactly one real run
    to style, even for "" (a deck that omits an optional field, e.g. no
    footer, must never crash the build). Returns the run."""
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text if text is not None else ""
    return run


def style_run(run, font_name, size, color_hex, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color_hex)


def add_footer(slide, theme, page_num, total, footer_text=""):
    if footer_text:
        _, tf = add_textbox(slide, Inches(0.5), SLIDE_H - Inches(0.45), Inches(8), Inches(0.35))
        style_run(set_text(tf, footer_text), theme["body_font"], 10, theme["secondary"])

    box2, tf2 = add_textbox(slide, SLIDE_W - Inches(1.6), SLIDE_H - Inches(0.45), Inches(1.2), Inches(0.35),
                          align=PP_ALIGN.RIGHT)
    box2.name = "ltr-text page"
    style_run(set_text(tf2, f"{page_num:02d} / {total:02d}"), theme["body_font"], 10, theme["secondary"])


def add_accent_bar(slide, theme, x=Inches(0.5), y=Inches(1.15), width=Inches(0.9), height=Pt(4)):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(theme["primary"])
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def add_line(slide, x, y, w, h, color_hex):
    """A thin filled rectangle used as a connector/divider line — simpler and
    more predictable across renderers than python-pptx's connector shapes."""
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(color_hex)
    line.line.fill.background()
    line.shadow.inherit = False
    return line


def add_circle(slide, cx, cy, diameter, fill_hex):
    d = diameter
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - d // 2, cy - d // 2, d, d)
    circ.fill.solid()
    circ.fill.fore_color.rgb = rgb(fill_hex)
    circ.line.fill.background()
    circ.shadow.inherit = False
    return circ


def add_panel(slide, x, y, w, h, fill_hex):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        panel.adjustments[0] = 0.06
    except Exception:
        pass
    panel.fill.solid()
    panel.fill.fore_color.rgb = rgb(fill_hex)
    panel.line.fill.background()
    panel.shadow.inherit = False
    return panel


def header(slide, theme, title, text_color):
    """Shared title + accent bar used by every content-slide layout, so
    every slide type lines up on the same grid."""
    _, tf = add_textbox(slide, Inches(0.7), Inches(0.5), Inches(11.9), Inches(0.9))
    style_run(set_text(tf, title), theme["heading_font"], 30, text_color, bold=True)
    add_accent_bar(slide, theme)


# ---------------------------------------------------------------- slide types

def slide_title(prs, theme, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    set_background(slide, theme, "hero")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]

    if spec.get("kicker"):
        _, tf = add_textbox(slide, Inches(0.9), Inches(1.75), Inches(11.5), Inches(0.42))
        style_run(set_text(tf, spec["kicker"].upper()), theme["body_font"], 14, theme["primary"], bold=True)

    # 1.7in comfortably fits a 44pt title up to 2 wrapped lines without the
    # box reaching into the subtitle's box below (validate_deck.py checks this).
    _, tf = add_textbox(slide, Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.7))
    style_run(set_text(tf, spec["title"]), theme["heading_font"], 44, text_color, bold=True)

    if spec.get("subtitle"):
        _, tf = add_textbox(slide, Inches(0.9), Inches(4.35), Inches(10.5), Inches(1.1))
        style_run(set_text(tf, spec["subtitle"]), theme["body_font"], 20, theme["secondary"])

    add_accent_bar(slide, theme, y=Inches(2.28))
    return slide


def slide_section(prs, theme, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "hero")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]

    # 72pt digits need ~1.4in of real box height (glyph + leading + margins).
    _, tf = add_textbox(slide, Inches(0.9), Inches(2.35), Inches(3), Inches(1.45))
    style_run(set_text(tf, spec.get("number", "")), theme["heading_font"], 72, theme["primary"], bold=True)

    _, tf = add_textbox(slide, Inches(0.9), Inches(3.95), Inches(10.5), Inches(1.3))
    style_run(set_text(tf, spec["title"]), theme["heading_font"], 36, text_color, bold=True)
    return slide


def slide_bullets(prs, theme, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)

    _, tf = add_textbox(slide, Inches(0.9), Inches(1.7), Inches(11.5), Inches(5.2))
    bullets = spec.get("bullets", [])
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(14)
        if isinstance(bullet, dict):
            head, detail = bullet.get("point", ""), bullet.get("detail", "")
        else:
            head, detail = bullet, ""
        run = p.add_run()
        run.text = f"●  {head}"
        style_run(run, theme["body_font"], 19, text_color, bold=bool(detail))
        if detail:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            p2.space_after = Pt(10)
            r2 = p2.add_run()
            r2.text = f"     {detail}"
            style_run(r2, theme["body_font"], 15, theme["secondary"])
    return slide


def slide_stats(prs, theme, spec):
    """Row of big KPI numbers — common in investor / QBR decks."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)

    stats = spec.get("stats", [])
    n = max(len(stats), 1)
    col_w = Inches(11.5) / n
    for i, s in enumerate(stats):
        x = Inches(0.9) + col_w * i
        _, tf = add_textbox(slide, x, Inches(2.4), col_w - Inches(0.3), Inches(1.4))
        style_run(set_text(tf, str(s.get("value", ""))), theme["heading_font"], 40, theme["primary"], bold=True)

        _, tf2 = add_textbox(slide, x, Inches(3.8), col_w - Inches(0.3), Inches(1.0))
        style_run(set_text(tf2, str(s.get("label", ""))), theme["body_font"], 16, theme["secondary"])
    return slide


def slide_quote(prs, theme, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]

    _, tf = add_textbox(slide, Inches(1.5), Inches(2.2), Inches(10.3), Inches(2.4))
    style_run(set_text(tf, f"“{spec['quote']}”"), theme["heading_font"], 28, text_color, italic=True)

    if spec.get("attribution"):
        _, tf2 = add_textbox(slide, Inches(1.5), Inches(4.8), Inches(10.3), Inches(0.6))
        style_run(set_text(tf2, f"— {spec['attribution']}"), theme["body_font"], 16, theme["primary"])
    return slide


def slide_table(prs, theme, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)

    headers_ = spec["headers"]
    rows = spec["rows"]
    n_rows, n_cols = len(rows) + 1, len(headers_)
    table_shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.9), Inches(1.8), Inches(11.5), Inches(0.5) * n_rows
    )
    table = table_shape.table
    for c, h in enumerate(headers_):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(theme["primary"])
        style_run(set_text(cell.text_frame, str(h), align=PP_ALIGN.LEFT), theme["body_font"], 14,
                   theme["bg"] if theme["style"] == "dark" else "FFFFFF", bold=True)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(theme["bg_alt"] if r % 2 == 0 else theme["bg"])
            style_run(set_text(cell.text_frame, str(val), align=PP_ALIGN.LEFT), theme["body_font"], 13, text_color)
    return slide


def slide_closing(prs, theme, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "hero")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]

    _, tf = add_textbox(slide, Inches(0.9), Inches(2.9), Inches(11.5), Inches(1.2))
    style_run(set_text(tf, spec.get("title", "Thank you")), theme["heading_font"], 40, text_color, bold=True)

    if spec.get("contact"):
        _, tf2 = add_textbox(slide, Inches(0.9), Inches(4.1), Inches(10.5), Inches(0.8))
        style_run(set_text(tf2, spec["contact"]), theme["body_font"], 18, theme["primary"])
    add_accent_bar(slide, theme, y=Inches(2.65))
    return slide


def slide_agenda(prs, theme, spec):
    """A numbered list with more air than `bullets` — for a talk outline or
    table of contents, not for dense supporting detail."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec.get("title", "Agenda"), text_color)

    items = spec.get("items", [])
    n = max(len(items), 1)
    top, bottom = 1.85, 6.9
    row_h = (bottom - top) / n
    font_size = 20 if n <= 5 else (16 if n <= 8 else 13)
    for i, item in enumerate(items):
        y = Inches(top + i * row_h)
        _, tf_num = add_textbox(slide, Inches(0.9), y, Inches(1.1), Inches(row_h - 0.08),
                                 anchor=MSO_ANCHOR.MIDDLE)
        style_run(set_text(tf_num, f"{i+1:02d}"), theme["heading_font"], font_size + 4, theme["primary"], bold=True)

        _, tf_item = add_textbox(slide, Inches(2.15), y, Inches(10.0), Inches(row_h - 0.08),
                                  anchor=MSO_ANCHOR.MIDDLE)
        style_run(set_text(tf_item, str(item)), theme["body_font"], font_size, text_color)

        if i < n - 1:
            add_line(slide, Inches(0.9), Inches(top + (i + 1) * row_h), Inches(10.5), Pt(0.75), theme["secondary"])
    return slide


def slide_two_column(prs, theme, spec):
    """Two independent lists side by side — problem/solution, before/after
    in plain text, current-state/future-state, etc."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)

    def column(x, heading, items):
        cursor = 1.9
        if heading:
            _, tf = add_textbox(slide, Inches(x), Inches(cursor), Inches(5.3), Inches(0.5))
            style_run(set_text(tf, heading), theme["heading_font"], 18, theme["primary"], bold=True)
            cursor += 0.65
        _, tf = add_textbox(slide, Inches(x), Inches(cursor), Inches(5.3), Inches(6.9 - cursor))
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(12)
            r = p.add_run()
            r.text = f"—  {it}"
            style_run(r, theme["body_font"], 16, text_color)

    column(0.9, spec.get("left_heading", ""), spec.get("left_items", []))
    column(6.9, spec.get("right_heading", ""), spec.get("right_items", []))
    add_line(slide, Inches(6.6), Inches(1.9), Pt(1), Inches(4.9), theme["secondary"])
    return slide


def slide_comparison(prs, theme, spec):
    """Two labeled cards for a direct A/B — before/after, us/them, current
    plan/proposed plan. More visually distinct than two_column's plain split."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)

    card_y, card_h, card_w, gap = 1.85, 4.95, 5.6, 0.33
    left_x, right_x = 0.7, 0.7 + card_w + gap

    def card(x, label, points, header_fill, header_text_color, body_fill):
        add_panel(slide, Inches(x), Inches(card_y), Inches(card_w), Inches(card_h), body_fill)
        strip = add_panel(slide, Inches(x), Inches(card_y), Inches(card_w), Inches(0.7), header_fill)
        _, tf = add_textbox(slide, Inches(x + 0.3), Inches(card_y), Inches(card_w - 0.6), Inches(0.7),
                             anchor=MSO_ANCHOR.MIDDLE)
        style_run(set_text(tf, label.upper()), theme["body_font"], 16, header_text_color, bold=True)
        _, tf2 = add_textbox(slide, Inches(x + 0.3), Inches(card_y + 0.9), Inches(card_w - 0.6),
                              Inches(card_h - 1.1))
        for i, pt in enumerate(points):
            p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(10)
            r = p.add_run()
            r.text = f"•  {pt}"
            style_run(r, theme["body_font"], 15, text_color)

    header_text_on_primary = theme["bg"] if theme["style"] == "dark" else "FFFFFF"
    card(left_x, spec.get("left_label", "Before"), spec.get("left_points", []),
         theme["secondary"], theme["bg"] if theme["style"] == "dark" else "FFFFFF", theme["bg_alt"])
    card(right_x, spec.get("right_label", "After"), spec.get("right_points", []),
         theme["primary"], header_text_on_primary, theme["bg_alt"])
    return slide


def slide_timeline(prs, theme, spec):
    """A horizontal chronological line with milestones — roadmap, history,
    launch schedule."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)

    milestones = spec.get("milestones", [])
    n = max(len(milestones), 1)
    line_y, x0, x1 = 4.2, 1.3, 12.0
    add_line(slide, Inches(x0), Inches(line_y), Inches(x1 - x0), Pt(2.5), theme["primary"])

    span = (x1 - x0) / max(n - 1, 1) if n > 1 else 0
    box_w = min(2.3, span if n > 1 else 4.0)
    for i, m in enumerate(milestones):
        cx = x0 + (span * i if n > 1 else (x1 - x0) / 2)
        add_circle(slide, Inches(cx), Inches(line_y + 0.0125), Inches(0.22), theme["primary"])
        _, tf_l = add_textbox(slide, Inches(cx - box_w / 2), Inches(line_y - 1.0), Inches(box_w), Inches(0.85),
                               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        style_run(set_text(tf_l, str(m.get("label", "")), align=PP_ALIGN.CENTER),
                  theme["heading_font"], 15, theme["primary"], bold=True)
        _, tf_t = add_textbox(slide, Inches(cx - box_w / 2), Inches(line_y + 0.3), Inches(box_w), Inches(1.4),
                               align=PP_ALIGN.CENTER)
        style_run(set_text(tf_t, str(m.get("text", "")), align=PP_ALIGN.CENTER), theme["body_font"], 13, text_color)
    return slide


def slide_process(prs, theme, spec):
    """Numbered circles for a how-it-works / step-by-step flow — distinct
    from `timeline` (which is chronological/dated) by reading as ordered
    steps rather than a calendar."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)

    steps = spec.get("steps", [])
    n = max(len(steps), 1)
    y_center, diameter, x0, x1 = 3.0, 1.05, 1.6, 11.7
    span = (x1 - x0) / max(n - 1, 1) if n > 1 else 0
    if n > 1:
        add_line(slide, Inches(x0), Inches(y_center - 0.02), Inches(x1 - x0), Pt(2), theme["secondary"])

    box_w = min(2.6, max(span, 1.8) if n > 1 else 4.0)
    circle_text_color = theme["bg"] if theme["style"] == "dark" else "FFFFFF"
    for i, s in enumerate(steps):
        cx = x0 + (span * i if n > 1 else (x1 - x0) / 2)
        if s.get("icon"):
            deck_extras.icon_or_glyph(slide, theme, s["icon"], str(i + 1), cx - diameter / 2, y_center - diameter / 2, diameter)
        else:
            add_circle(slide, Inches(cx), Inches(y_center), Inches(diameter), theme["primary"])
            _, tf_n = add_textbox(slide, Inches(cx - diameter / 2), Inches(y_center - diameter / 2),
                                   Inches(diameter), Inches(diameter), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            style_run(set_text(tf_n, str(s.get("number", i + 1)), align=PP_ALIGN.CENTER),
                      theme["heading_font"], 26, circle_text_color, bold=True)

        _, tf_l = add_textbox(slide, Inches(cx - box_w / 2), Inches(y_center + diameter / 2 + 0.25),
                               Inches(box_w), Inches(0.8), align=PP_ALIGN.CENTER)
        style_run(set_text(tf_l, str(s.get("label", "")), align=PP_ALIGN.CENTER),
                  theme["heading_font"], 16, text_color, bold=True)
        _, tf_d = add_textbox(slide, Inches(cx - box_w / 2), Inches(y_center + diameter / 2 + 1.1),
                               Inches(box_w), Inches(1.6), align=PP_ALIGN.CENTER)
        style_run(set_text(tf_d, str(s.get("detail", "")), align=PP_ALIGN.CENTER), theme["body_font"], 13,
                  theme["secondary"])
    return slide


def slide_big_statement(prs, theme, spec):
    """One huge, unavoidable line — a milestone number, a mission statement,
    a single hard-hitting claim. Use sparingly, once or twice a deck."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "hero")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]

    statement = spec["statement"]
    size = 54 if len(statement) < 40 else (40 if len(statement) < 80 else 30)
    if deck_extras.RTL:
        size -= 8 if size > 30 else 4       # Arabic glyphs are taller at the same size
    box_h = 2.6 if size >= 40 else 3.2
    _, tf = add_textbox(slide, Inches(1.15), Inches(2.6), Inches(11.0), Inches(box_h),
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    style_run(set_text(tf, statement, align=PP_ALIGN.CENTER), theme["heading_font"], size, text_color, bold=True)

    if spec.get("support"):
        _, tf2 = add_textbox(slide, Inches(1.65), Inches(2.6 + box_h + 0.1), Inches(10.0), Inches(0.9),
                              align=PP_ALIGN.CENTER)
        style_run(set_text(tf2, spec["support"], align=PP_ALIGN.CENTER), theme["body_font"], 18, theme["primary"])
    return slide


def slide_team(prs, theme, spec):
    """A grid of people — initials in a colored circle stand in for headshots
    since this skill never fabricates photos of real people."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec.get("title", "Team"), text_color)

    people = spec.get("people", [])
    n = max(len(people), 1)
    cols = min(n, 5)
    rows_n = math.ceil(n / cols)
    card_w = 11.4 / cols
    row_h = min(2.7, (6.9 - 1.9) / rows_n)
    avatar_d = 1.0
    circle_text_color = theme["bg"] if theme["style"] == "dark" else "FFFFFF"

    for i, person in enumerate(people):
        col, row = i % cols, i // cols
        cx = 0.9 + card_w * col + card_w / 2
        top = 2.0 + row * row_h
        initials = "".join(w[0].upper() for w in str(person.get("name", "?")).split()[:2]) or "?"
        photo = deck_extras.asset(person.get("photo")) if person.get("photo") else None
        if photo:   # the person's OWN photo (a file the user gave) — never a stock face for a real person
            pic = deck_extras.picture_cover(slide, photo, cx - avatar_d / 2, top, avatar_d, avatar_d)
            pic.auto_shape_type = MSO_SHAPE.OVAL
        else:
            add_circle(slide, Inches(cx), Inches(top + avatar_d / 2), Inches(avatar_d), theme["primary"])
            _, tf_i = add_textbox(slide, Inches(cx - avatar_d / 2), Inches(top), Inches(avatar_d), Inches(avatar_d),
                                   align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            style_run(set_text(tf_i, initials, align=PP_ALIGN.CENTER), theme["heading_font"], 22, circle_text_color,
                      bold=True)

        _, tf_n = add_textbox(slide, Inches(cx - card_w / 2 + 0.1), Inches(top + avatar_d + 0.15),
                               Inches(card_w - 0.2), Inches(0.45), align=PP_ALIGN.CENTER)
        style_run(set_text(tf_n, str(person.get("name", "")), align=PP_ALIGN.CENTER),
                  theme["heading_font"], 15, text_color, bold=True)
        _, tf_r = add_textbox(slide, Inches(cx - card_w / 2 + 0.1), Inches(top + avatar_d + 0.58),
                               Inches(card_w - 0.2), Inches(0.5), align=PP_ALIGN.CENTER)
        style_run(set_text(tf_r, str(person.get("role", "")), align=PP_ALIGN.CENTER), theme["body_font"], 12,
                  theme["secondary"])
    return slide


def slide_feature(prs, theme, spec):
    """Text + a stat callout panel on the side — the layout to reach for
    when a slide "needs an image": a real, meaningful number in a styled
    panel reads as more premium than a stock photo, and this skill never
    fabricates a fake photo to fill space."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)

    cursor = 1.9
    if spec.get("description"):
        _, tf = add_textbox(slide, Inches(0.9), Inches(cursor), Inches(6.3), Inches(1.55))
        style_run(set_text(tf, spec["description"]), theme["body_font"], 17, text_color)
        cursor += 1.75

    _, tf = add_textbox(slide, Inches(0.9), Inches(cursor), Inches(6.3), Inches(6.8 - cursor))
    for i, b in enumerate(spec.get("bullets", [])):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(12)
        r = p.add_run()
        r.text = f"●  {b}"
        style_run(r, theme["body_font"], 16, text_color)

    panel_x, panel_w = 7.8, 4.6
    img = deck_extras.asset(spec.get("image"), size=(1100, 1170)) if spec.get("image") else None
    if img:
        deck_extras.picture_cover(slide, img, panel_x, 1.9, panel_w, 4.9, rounded=0.05)
        if spec.get("panel_value"):
            deck_extras.rect(slide, panel_x + 0.3, 5.1, panel_w - 0.6, 1.45, theme["bg"], alpha=0.9,
                             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
            _, tfv = add_textbox(slide, Inches(panel_x + 0.45), Inches(5.12), Inches(panel_w - 0.9), Inches(0.8),
                                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            style_run(set_text(tfv, str(spec["panel_value"]), align=PP_ALIGN.CENTER), theme["heading_font"], 30,
                      theme["primary"], bold=True)
            _, tfc = add_textbox(slide, Inches(panel_x + 0.45), Inches(5.9), Inches(panel_w - 0.9), Inches(0.55),
                                  align=PP_ALIGN.CENTER)
            style_run(set_text(tfc, str(spec.get("panel_caption") or spec.get("panel_label") or ""), align=PP_ALIGN.CENTER),
                      theme["body_font"], 12, text_color)
        return slide
    add_panel(slide, Inches(panel_x), Inches(1.9), Inches(panel_w), Inches(4.9), theme["bg_alt"])
    py = 2.5
    if spec.get("panel_label"):
        _, tfl = add_textbox(slide, Inches(panel_x + 0.4), Inches(py), Inches(panel_w - 0.8), Inches(0.4),
                              align=PP_ALIGN.CENTER)
        style_run(set_text(tfl, spec["panel_label"].upper(), align=PP_ALIGN.CENTER), theme["body_font"], 13,
                  theme["secondary"], bold=True)
        py += 0.6
    _, tfv = add_textbox(slide, Inches(panel_x + 0.3), Inches(py), Inches(panel_w - 0.6), Inches(1.6),
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    style_run(set_text(tfv, str(spec.get("panel_value", "")), align=PP_ALIGN.CENTER), theme["heading_font"], 36,
              theme["primary"], bold=True)
    py += 1.7
    if spec.get("panel_caption"):
        _, tfc = add_textbox(slide, Inches(panel_x + 0.4), Inches(py), Inches(panel_w - 0.8), Inches(0.8),
                              align=PP_ALIGN.CENTER)
        style_run(set_text(tfc, spec["panel_caption"], align=PP_ALIGN.CENTER), theme["body_font"], 14, text_color)
    return slide


def slide_icon_grid(prs, theme, spec):
    """A grid of feature callouts — each item gets a colored glyph circle,
    a bold label, and a short description. For "what we offer" / "key
    capabilities" slides where `bullets` would read as too flat."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)

    items = spec.get("items", [])
    n = max(len(items), 1)
    cols = min(n, 3)
    rows_n = math.ceil(n / cols)
    gap = 0.4
    col_w = (11.6 - gap * (cols - 1)) / cols
    row_h = min(2.35, (6.85 - 1.85) / rows_n)
    glyph_d = 0.6
    circle_text_color = theme["bg"] if theme["style"] == "dark" else "FFFFFF"

    for i, item in enumerate(items):
        col, row = i % cols, i // cols
        x = 0.85 + col * (col_w + gap)
        top = 1.9 + row * row_h
        glyph = str(item.get("glyph", "")) or str(i + 1)
        if item.get("icon"):
            deck_extras.icon_or_glyph(slide, theme, item["icon"], glyph, x, top, glyph_d)
        else:
            add_circle(slide, Inches(x + glyph_d / 2), Inches(top + glyph_d / 2), Inches(glyph_d), theme["primary"])
            _, tf_g = add_textbox(slide, Inches(x), Inches(top), Inches(glyph_d), Inches(glyph_d),
                                   align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            style_run(set_text(tf_g, glyph[:2].upper(), align=PP_ALIGN.CENTER), theme["heading_font"], 13,
                      circle_text_color, bold=True)

        _, tf_l = add_textbox(slide, Inches(x), Inches(top + glyph_d + 0.15), Inches(col_w), Inches(0.4))
        style_run(set_text(tf_l, item.get("label", "")), theme["heading_font"], 16, theme["primary"], bold=True)

        _, tf_d = add_textbox(slide, Inches(x), Inches(top + glyph_d + 0.6), Inches(col_w),
                               Inches(row_h - glyph_d - 0.7))
        style_run(set_text(tf_d, item.get("description", "")), theme["body_font"], 13, text_color)
    return slide


def slide_pricing(prs, theme, spec):
    """2-4 pricing-tier cards side by side, one optionally highlighted as
    the recommended plan."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)

    tiers = spec.get("tiers", [])
    n = max(len(tiers), 1)
    gap = 0.3
    card_y, card_h = 1.95, 4.85
    card_w = (11.6 - gap * (n - 1)) / n
    header_on_primary = theme["bg"] if theme["style"] == "dark" else "FFFFFF"

    for i, tier in enumerate(tiers):
        x = 0.85 + i * (card_w + gap)
        highlight = bool(tier.get("highlight"))
        head_fill = theme["primary"] if highlight else theme["bg_alt"]
        head_text = header_on_primary if highlight else theme["primary"]
        add_panel(slide, Inches(x), Inches(card_y), Inches(card_w), Inches(card_h), theme["bg_alt"])
        strip = add_panel(slide, Inches(x), Inches(card_y), Inches(card_w), Inches(1.15), head_fill)
        _, tf_n = add_textbox(slide, Inches(x + 0.2), Inches(card_y + 0.12), Inches(card_w - 0.4), Inches(0.35),
                               align=PP_ALIGN.CENTER)
        style_run(set_text(tf_n, str(tier.get("name", "")).upper(), align=PP_ALIGN.CENTER), theme["body_font"], 13,
                  head_text, bold=True)
        _, tf_p = add_textbox(slide, Inches(x + 0.15), Inches(card_y + 0.45), Inches(card_w - 0.3), Inches(0.65),
                               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        price_line = str(tier.get("price", ""))
        if tier.get("period"):
            price_line += f" / {tier['period']}"
        style_run(set_text(tf_p, price_line, align=PP_ALIGN.CENTER), theme["heading_font"], 22, head_text, bold=True)

        _, tf_f = add_textbox(slide, Inches(x + 0.25), Inches(card_y + 1.35), Inches(card_w - 0.5),
                               Inches(card_h - 1.55))
        for j, feat in enumerate(tier.get("features", [])):
            p = tf_f.paragraphs[0] if j == 0 else tf_f.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(9)
            r = p.add_run()
            r.text = f"✓  {feat}"
            style_run(r, theme["body_font"], 13, text_color)
    return slide


def slide_faq(prs, theme, spec):
    """A short Q&A list — each question bold, answer beneath it in a
    quieter color. Keep to ~4 pairs; validate_deck.py will flag overflow
    if the copy runs long."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec.get("title", "FAQ"), text_color)

    items = spec.get("items", [])
    n = max(len(items), 1)
    top, bottom = 1.9, 6.9
    row_h = (bottom - top) / n
    for i, item in enumerate(items):
        y = top + i * row_h
        _, tf_q = add_textbox(slide, Inches(0.9), Inches(y), Inches(10.8), Inches(row_h * 0.48))
        style_run(set_text(tf_q, item.get("question", "")), theme["heading_font"], 16, theme["primary"], bold=True)
        _, tf_a = add_textbox(slide, Inches(0.9), Inches(y + row_h * 0.46), Inches(10.8), Inches(row_h * 0.46))
        style_run(set_text(tf_a, item.get("answer", "")), theme["body_font"], 13, text_color)
        if i < n - 1:
            add_line(slide, Inches(0.9), Inches(y + row_h - 0.08), Inches(10.8), Pt(0.75), theme["secondary"])
    return slide


def slide_chart_bar(prs, theme, spec):
    """A real, native, editable bar chart drawn from rectangles (not a
    picture) — every bar's height and label is an actual pptx shape you
    can click and adjust."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)

    bars = spec.get("bars", [])
    n = max(len(bars), 1)
    x0, x1 = 1.2, 12.1
    baseline_y, top_y = 6.3, 2.1
    max_h = baseline_y - top_y
    values = [float(b.get("value", 0)) for b in bars] or [0]
    vmax = max(values) or 1

    add_line(slide, Inches(x0), Inches(baseline_y), Inches(x1 - x0), Pt(1.5), theme["secondary"])
    span = (x1 - x0) / n
    bar_w = min(1.1, span * 0.5)
    for i, b in enumerate(bars):
        val = float(b.get("value", 0))
        h = max(0.12, (val / vmax) * max_h)
        cx = x0 + span * (i + 0.5)
        add_panel(slide, Inches(cx - bar_w / 2), Inches(baseline_y - h), Inches(bar_w), Inches(h), theme["primary"])
        _, tf_v = add_textbox(slide, Inches(cx - span / 2), Inches(baseline_y - h - 0.48), Inches(span), Inches(0.42),
                               align=PP_ALIGN.CENTER)
        style_run(set_text(tf_v, str(b.get("display", b.get("value", ""))), align=PP_ALIGN.CENTER),
                  theme["heading_font"], 14, theme["primary"], bold=True)
        _, tf_l = add_textbox(slide, Inches(cx - span / 2), Inches(baseline_y + 0.1), Inches(span), Inches(0.65),
                               align=PP_ALIGN.CENTER)
        style_run(set_text(tf_l, str(b.get("label", "")), align=PP_ALIGN.CENTER), theme["body_font"], 12, text_color)
    return slide


def slide_logo_wall(prs, theme, spec):
    """A grid of partner/client names as styled text badges — this skill
    never fabricates a company's actual logo artwork, so a clean
    typographic badge stands in for a logo wall."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec.get("title", "Trusted by"), text_color)

    names = spec.get("logos", [])
    n = max(len(names), 1)
    cols = min(n, 4)
    rows_n = math.ceil(n / cols)
    gap = 0.35
    badge_w = (11.6 - gap * (cols - 1)) / cols
    badge_h = min(1.3, (6.7 - 1.9 - gap * (rows_n - 1)) / rows_n)

    for i, name in enumerate(names):
        col, row = i % cols, i // cols
        x = 0.85 + col * (badge_w + gap)
        y = 1.9 + row * (badge_h + gap)
        add_panel(slide, Inches(x), Inches(y), Inches(badge_w), Inches(badge_h), theme["bg_alt"])
        _, tf = add_textbox(slide, Inches(x + 0.15), Inches(y), Inches(badge_w - 0.3), Inches(badge_h),
                             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        style_run(set_text(tf, str(name), align=PP_ALIGN.CENTER), theme["heading_font"], 15, theme["primary"],
                  bold=True)
    return slide


def slide_vertical_timeline(prs, theme, spec):
    """A top-to-bottom milestone list — the variant to reach for when a
    horizontal `timeline` would be too cramped (more than ~4 milestones,
    or entries with real supporting detail rather than one short line)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec.get("title", "Timeline"), text_color)

    milestones = spec.get("milestones", [])
    n = max(len(milestones), 1)
    top, bottom, line_x = 2.05, 6.85, 1.35
    row_h = (bottom - top) / n
    if n > 1:
        add_line(slide, Inches(line_x - 0.01), Inches(top + row_h / 2), Pt(2),
                 Inches(row_h * (n - 1)), theme["primary"])
    for i, m in enumerate(milestones):
        cy = top + row_h * (i + 0.5)
        add_circle(slide, Inches(line_x), Inches(cy), Inches(0.22), theme["primary"])
        _, tf_l = add_textbox(slide, Inches(1.75), Inches(cy - row_h / 2 + 0.05), Inches(2.2), Inches(row_h - 0.1),
                               anchor=MSO_ANCHOR.MIDDLE)
        style_run(set_text(tf_l, str(m.get("label", ""))), theme["heading_font"], 15, theme["primary"], bold=True)
        _, tf_t = add_textbox(slide, Inches(4.1), Inches(cy - row_h / 2 + 0.05), Inches(8.3), Inches(row_h - 0.1),
                               anchor=MSO_ANCHOR.MIDDLE)
        style_run(set_text(tf_t, str(m.get("text", ""))), theme["body_font"], 14, text_color)
    return slide


def slide_array_pass(prs, theme, spec):
    """A binary-search picture: one row of array cells per step.

    Cells inside the current window stay bright. The middle cell uses the
    accent color. Cells outside the window are dim. L / M / R sit under the
    left, mid, and right indexes. Text is real PowerPoint text, not a picture.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    on_accent = theme["bg"] if theme["style"] == "dark" else "FFFFFF"
    header(slide, theme, spec["title"], text_color)

    steps = spec.get("steps") or [spec]
    n_rows = max(len(steps), 1)
    top, bottom = 1.7, 6.85
    row_h = (bottom - top) / n_rows
    dim = "101828" if theme["style"] == "dark" else "E4E7EE"

    for r, step in enumerate(steps):
        values = [str(v) for v in step.get("values", [])]
        n = max(len(values), 1)
        left = int(step.get("left", 0))
        right = int(step.get("right", n - 1))
        mid = step.get("mid", None)
        mid = None if mid is None else int(mid)
        caption = str(step.get("caption", ""))

        margin, gap = 0.7, 0.1
        usable = 11.95
        cell_w = (usable - gap * (n - 1)) / n
        cell_h = 1.35 if n_rows == 1 else 0.42
        y = top + r * row_h
        value_size = 28 if n_rows == 1 else 16
        index_size = 12

        for i, val in enumerate(values):
            x = margin + i * (cell_w + gap)
            in_window = left <= i <= right
            is_mid = mid is not None and i == mid and in_window
            if is_mid:
                fill, color = theme["primary"], on_accent
            elif in_window:
                fill, color = theme["bg_alt"], text_color
            else:
                fill, color = dim, theme["secondary"]
            add_panel(slide, Inches(x), Inches(y), Inches(cell_w), Inches(cell_h), fill)
            _, tf_v = add_textbox(
                slide, Inches(x), Inches(y), Inches(cell_w), Inches(cell_h),
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
            )
            style_run(set_text(tf_v, val, align=PP_ALIGN.CENTER), theme["heading_font"], value_size, color, bold=True)

            _, tf_i = add_textbox(
                slide, Inches(x), Inches(y + cell_h + 0.02), Inches(cell_w), Inches(0.40),
                align=PP_ALIGN.CENTER,
            )
            style_run(set_text(tf_i, str(i), align=PP_ALIGN.CENTER), theme["body_font"], index_size, theme["secondary"])

            roles = []
            if in_window and i == left:
                roles.append("L")
            if in_window and mid is not None and i == mid:
                roles.append("M")
            if in_window and i == right and "R" not in roles:
                roles.append("R")
            if roles:
                _, tf_role = add_textbox(
                    slide, Inches(x), Inches(y + cell_h + 0.44), Inches(cell_w), Inches(0.40),
                    align=PP_ALIGN.CENTER,
                )
                style_run(
                    set_text(tf_role, " ".join(roles), align=PP_ALIGN.CENTER),
                    theme["heading_font"], 12, theme["primary"], bold=True,
                )

        cap_y = y + cell_h + 0.88
        cap_h = row_h - (cell_h + 0.92)
        _, tf_c = add_textbox(slide, Inches(0.7), Inches(cap_y), Inches(11.95), Inches(cap_h))
        style_run(set_text(tf_c, caption), theme["body_font"], 16 if n_rows == 1 else 13, text_color)
    return slide


def _fill_lines(tf, lines, font_name, size, color_hex):
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line if line else " "
        style_run(run, font_name, size, color_hex)


def _token_style(ttype):
    """IDE colours: blue keywords, gold strings, green names, pink tags, grey comments."""
    name = "" if ttype is None else str(ttype)
    if "Comment" in name:
        return "8B9AAB", True
    if "Keyword" in name or "Name.Builtin" in name:
        return "7EE0FF", False
    if "String" in name:
        return "E8C547", False
    if "Number" in name:
        return "C3A6FF", False
    if any(k in name for k in ("Name.Function", "Name.Class", "Name.Namespace")):
        return "9BE36D", False
    if "Name.Tag" in name or name.endswith("Operator") or ".Operator" in name:
        return "FF6B9D", False
    return "E6EDF3", False


def _lexer_for(code, lang):
    from pygments.lexers import PhpLexer, SqlLexer, get_lexer_by_name

    if lang:
        return get_lexer_by_name(lang, stripall=False)
    head = code.lstrip()
    if head.startswith(("CREATE", "INSERT", "SELECT", "UPDATE", "DELETE")):
        return SqlLexer()
    if "<form" in code or "<html" in code or "<div" in code:
        return get_lexer_by_name("html+php")
    return PhpLexer(startinline=True)


def _paint_code(tf, code, lang=None, size=16):
    src = code.replace("\t", "    ").strip("\n") + "\n"
    try:
        tokens = list(_lexer_for(src, lang).get_tokens(src))
    except Exception:
        tokens = [(None, src)]
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    for ttype, value in tokens:
        color, italic = _token_style(ttype)
        parts = value.split("\n")
        for i, part in enumerate(parts):
            if i > 0:
                p = tf.add_paragraph()
                p.alignment = PP_ALIGN.LEFT
                p.space_before = Pt(0)
                p.space_after = Pt(0)
            if part == "":
                if i > 0:
                    run = p.add_run()
                    run.text = " "
                    style_run(run, "Consolas", size, "E6EDF3")
                continue
            run = p.add_run()
            run.text = part
            style_run(run, "Consolas", size, color, italic=italic)


def slide_code(prs, theme, spec):
    """A dark editor panel. Tokens are coloured; every run is still editable text."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)
    code = str(spec.get("code", ""))
    nlines = max(code.strip("\n").count("\n") + 1, 1)
    panel_h = min(4.85, 0.72 + 0.30 * nlines)
    add_panel(slide, Inches(0.55), Inches(1.7), Inches(12.2), Inches(panel_h), "1E2433").name = "keep-ltr panel"
    for i, col in enumerate(("FF5F57", "FEBC2E", "28C840")):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.78 + i * 0.22), Inches(1.86), Inches(0.12), Inches(0.12)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = rgb(col)
        dot.line.fill.background()
        dot.shadow.inherit = False
        dot.name = "keep-ltr dot"
    box, tf = add_textbox(slide, Inches(0.78), Inches(2.08), Inches(11.7), Inches(max(0.45, panel_h - 0.52)))
    box.name = "keep-ltr code"
    tf.word_wrap = False
    _paint_code(tf, code, spec.get("lang"))
    if spec.get("caption"):
        _, tfc = add_textbox(slide, Inches(0.7), Inches(min(6.3, 1.7 + panel_h + 0.25)), Inches(11.9), Inches(0.75))
        style_run(set_text(tfc, spec["caption"]), theme["body_font"], 17, text_color)
    return slide


def slide_shot(prs, theme, spec):
    """A real screenshot, kept as a picture. The caption under it is editable text."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)
    src = spec.get("src", "")
    if spec.get("image"):
        src = deck_extras.asset(spec["image"]) or src
    cap_y = 6.45
    if src and Path(src).exists():
        from PIL import Image
        iw, ih = Image.open(src).size
        aspect = iw / ih if ih else 1
        box_w, box_h = 11.6, 4.55
        if box_w / box_h > aspect:
            h, w = box_h, box_h * aspect
        else:
            w, h = box_w, box_w / aspect
        x = 0.85 + (11.6 - w) / 2
        slide.shapes.add_picture(str(src), Inches(x), Inches(1.6), Inches(w), Inches(h))
        cap_y = 1.6 + h + 0.08
    _, tf = add_textbox(slide, Inches(0.7), Inches(cap_y), Inches(11.9), Inches(0.4))
    style_run(set_text(tf, spec.get("caption", "")), theme["body_font"], 14, theme["secondary"])
    return slide


def slide_relation(prs, theme, spec):
    """Entity cards with the cardinality written in the gap. Native shapes, not a bitmap."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, theme, "content")
    text_color = theme["text_on_dark"] if theme["style"] == "dark" else theme["text_on_light"]
    header(slide, theme, spec["title"], text_color)

    cards = spec.get("cards") or []
    n = max(len(cards), 1)
    gap = 1.15 if n == 2 else 0.7
    margin = 0.55
    usable = 13.333 - margin * 2 - gap * (n - 1)
    cw = usable / n
    top, ch = 1.85, 3.55
    for i, card in enumerate(cards):
        x = margin + i * (cw + gap)
        add_panel(slide, Inches(x), Inches(top), Inches(cw), Inches(ch), "161D32")
        _, tf_h = add_textbox(slide, Inches(x + 0.2), Inches(top + 0.16), Inches(cw - 0.4), Inches(0.55))
        style_run(set_text(tf_h, card.get("name", "")), theme["heading_font"], 20, "E0B56A", bold=True)
        _, tf_b = add_textbox(slide, Inches(x + 0.2), Inches(top + 0.82), Inches(cw - 0.4), Inches(2.45))
        _fill_lines(tf_b, card.get("lines") or [], theme["body_font"], 16, "F7F3EC")
        if i < n - 1:
            mark = (spec.get("marks") or ["", ""])[i] if i < len(spec.get("marks") or []) else ""
            gx = x + cw
            _, tf_m = add_textbox(slide, Inches(gx), Inches(top + 1.35), Inches(gap), Inches(0.7), align=PP_ALIGN.CENTER)
            style_run(set_text(tf_m, mark, align=PP_ALIGN.CENTER), theme["heading_font"], 18, theme["primary"], bold=True)
    note = spec.get("note", "")
    if note:
        _, tf_n = add_textbox(slide, Inches(0.7), Inches(5.6), Inches(11.9), Inches(0.7))
        style_run(set_text(tf_n, note), theme["body_font"], 16, text_color)
    return slide


import deck_extras  # noqa: E402  (v2 slide types, images, RTL — see deck_extras.py)

SLIDE_BUILDERS = {
    "title": slide_title,
    "section": slide_section,
    "bullets": slide_bullets,
    "stats": slide_stats,
    "quote": slide_quote,
    "table": slide_table,
    "closing": slide_closing,
    "agenda": slide_agenda,
    "two_column": slide_two_column,
    "comparison": slide_comparison,
    "timeline": slide_timeline,
    "process": slide_process,
    "big_statement": slide_big_statement,
    "team": slide_team,
    "feature": slide_feature,
    "icon_grid": slide_icon_grid,
    "pricing": slide_pricing,
    "faq": slide_faq,
    "chart_bar": slide_chart_bar,
    "logo_wall": slide_logo_wall,
    "vertical_timeline": slide_vertical_timeline,
    "array_pass": slide_array_pass,
    "code": slide_code,
    "shot": slide_shot,
    "relation": slide_relation,
}
SLIDE_BUILDERS.update(deck_extras.register(sys.modules[__name__]))


def build(deck_spec_path: str, out_path: str, *, offline: bool = False):
    spec_path = Path(deck_spec_path)
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    out_path = Path(out_path)
    from pdk.text import spec_is_rtl
    from pdk.themes import tokens, ARABIC_FONTS
    from pdk.assets import Assets

    theme = dict(get_theme(spec.get("theme", "plain-pro"), spec.get("brand")))
    rtl = spec_is_rtl(spec)
    if rtl:
        style = spec.get("arabic_style", "modern")
        h, b = ARABIC_FONTS.get(style, ARABIC_FONTS["modern"])
        theme["heading_font"] = theme.get("heading_font_ar") or spec.get("heading_font") or h
        theme["body_font"] = theme.get("body_font_ar") or spec.get("body_font") or b
    else:
        theme["heading_font"] = spec.get("heading_font") or theme["heading_font"]
        theme["body_font"] = spec.get("body_font") or theme["body_font"]
    tk = tokens(theme, arabic=rtl)
    deck_extras.RTL = rtl
    deck_extras.TOKENS = tk
    deck_extras.ASSETS = Assets(out_path.parent / f"_assets-{out_path.stem}", tk, base=spec_path.parent,
                                project=spec.get("project") or out_path.stem, offline=offline,
                                icon_set=spec.get("icon_set", "material-symbols"))
    _BG_CACHE.clear()

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slides_spec = list(spec["slides"])
    total = len(slides_spec)
    transition = spec.get("transition", "fade")
    built = []
    for i, s in enumerate(slides_spec, start=1):
        builder = SLIDE_BUILDERS.get(s["type"])
        if builder is None:
            raise ValueError(
                f"Unknown slide type '{s['type']}' (slide {i}). Available: {', '.join(SLIDE_BUILDERS)}"
            )
        slide = builder(prs, theme, s)
        if s.get("background"):
            deck_extras.apply_background(slide, theme, s["background"], s.get("overlay", 0.6),
                                         s.get("overlay_direction", "full"))
        if s["type"] not in ("title", "image_hero"):
            add_footer(slide, theme, i, total, footer_text=spec.get("footer", ""))
        add_transition(slide, kind=s.get("transition", transition))
        deck_extras.set_notes(slide, s.get("notes"))
        built.append(slide)

    credits = deck_extras.ASSETS.credits()
    if credits and spec.get("credits", True):
        slide = deck_extras.slide_credits(prs, theme, {"lines": credits})
        add_transition(slide, kind=transition)
        built.append(slide)
    if rtl:
        for slide in built:
            deck_extras.mirror_slide(slide)
    for slide in built:
        deck_extras.rtl_font_runs(slide, theme["heading_font"], theme["body_font"]) if rtl else None
    if spec.get("meta_title") or spec.get("title"):
        prs.core_properties.title = spec.get("meta_title") or spec.get("title")
    if spec.get("author"):
        prs.core_properties.author = spec["author"]

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print(f"Wrote {out_path}  ({len(built)} slides, theme={theme['id']}{', RTL' if rtl else ''})")
    for w in deck_extras.ASSETS.warnings:
        print("  WARN:", w)
    if credits:
        print(f"  {len(credits)} credit line(s) → last slide")
    return str(out_path)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 2:
        print("Usage: python3 build_deck.py deck_spec.json output.pptx [--offline]")
        sys.exit(1)
    build(args[0], args[1], offline="--offline" in sys.argv)
