#!/usr/bin/env python3
"""build_sheet.py — a workbook spec → a professional, formula-driven .xlsx.

    run.sh build_sheet.py sheet_spec.json out/model.xlsx [--no-recalc]

What you get: real Excel Tables (banded, filterable, structured), live formulas (never pasted numbers
for anything that is calculated), totals rows, named input cells, data validation drop-downs,
conditional formatting (data bars, colour scales, icon sets, rules), native editable charts, a KPI
dashboard sheet, a contents sheet with links, frozen headers, print setup, and right-to-left sheets
for Arabic. After saving, LibreOffice recalculates a copy and every formula error (#DIV/0!, #REF!,
#NAME?, #VALUE!, #N/A) is reported with its cell — fix them before delivering.

Spec: see ../references/sheet-spec.md.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
KIT = HERE.parents[1] / "premium-design-kit" / "lib"
if str(KIT) not in sys.path:
    sys.path.insert(0, str(KIT))

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, IconSetRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from pdk.text import spec_is_rtl
from pdk.themes import ARABIC_FONTS, mix, resolve_theme, tokens

FORMATS = {
    "text": "@", "int": "#,##0", "number": "#,##0.00", "decimal": "#,##0.00", "percent": "0.0%", "pct": "0%",
    "date": "yyyy-mm-dd", "datetime": "yyyy-mm-dd hh:mm", "time": "hh:mm", "duration": "[h]:mm",
    "money": '#,##0.00', "score": "0.0", "bool": "General",
}
ERRORS = ("#DIV/0!", "#REF!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "Err:")


class SheetBuilder:
    def __init__(self, spec: dict, base: Path):
        self.spec = spec
        self.base = base
        self.rtl = spec_is_rtl(spec)
        th = resolve_theme(spec.get("theme"), spec.get("brand"), default="plain-pro")
        self.tk = tokens(th, arabic=self.rtl)
        if self.rtl:
            self.font = th.get("body_font_ar") or spec.get("font") or "Arial"
        else:
            self.font = spec.get("font") or ("Calibri" if th["body_font"] in ("Helvetica", "Georgia") else th["body_font"])
        self.cur = spec.get("currency", "")
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.tables = {}      # sheet name → {col name: letter, "first": r, "last": r, "ref": ...}
        self.names = {}
        self.warnings = []

    # ------------------------------------------------------------------ styles
    def f(self, **kw):
        base = {"name": self.font, "size": 10.5, "color": self.tk["ink"]}
        base.update(kw)
        return Font(**base)

    def fill(self, hexc):
        return PatternFill("solid", start_color=hexc, end_color=hexc)

    @property
    def thin(self):
        s = Side(style="thin", color=self.tk["line"])
        return Border(bottom=s)

    def fmt(self, col):
        t = col.get("type", "text")
        if col.get("format"):
            return col["format"]
        if t == "money":
            c = col.get("currency", self.cur)
            if c and len(c) == 1:
                return f'"{c}"#,##0.00'
            return f'#,##0.00 "{c}"' if c else "#,##0.00"
        return FORMATS.get(t, "General")

    # ------------------------------------------------------------------ build
    def build(self, out: Path, recalc=True) -> dict:
        sheets = self.spec.get("sheets", [])
        if self.spec.get("contents", len(sheets) > 2):
            self.contents_sheet(sheets)
        for sh in sheets:
            kind = sh.get("kind", "table")
            ws = self.wb.create_sheet(sh["name"][:31])
            ws.sheet_view.showGridLines = sh.get("gridlines", kind == "table" and False)
            ws.sheet_view.rightToLeft = self.rtl
            if sh.get("tab_color", True):
                ws.sheet_properties.tabColor = self.tk["accent"] if kind == "dashboard" else mix(self.tk["accent"], "FFFFFF", 0.5)
            getattr(self, f"sheet_{kind}")(ws, sh)
            self.print_setup(ws, sh)
        # second pass: {Sheet.Column} references anywhere → real ranges; charts
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("=") and "{" in c.value:
                        c.value = self.resolve_refs(c.value)
        for sh in sheets:
            ws = self.wb[sh["name"][:31]]
            for ch in sh.get("charts", []):
                self.chart(ws, ch)
        props = self.wb.properties
        props.title = self.spec.get("title", "")
        props.creator = self.spec.get("author", "")
        if self.wb.sheetnames and self.spec.get("contents", len(sheets) > 2):
            self.wb.active = 1 if len(self.wb.sheetnames) > 1 else 0
        out.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(out)
        rep = {"file": str(out), "sheets": self.wb.sheetnames, "warnings": self.warnings}
        if recalc:
            rep.update(recalc_check(out))
        return rep

    def resolve_refs(self, f):
        """{Sheet.Column} → 'Sheet'!$D$5:$D$9 · {Sheet.Column.last} / .first → that single cell ·
        {Sheet.Column.total} → the totals-row cell."""
        def rep(m):
            parts = m.group(1).split(".")
            if len(parts) < 2 or parts[0] not in self.tables:
                return m.group(0)
            t = self.tables[parts[0]]
            L = t["letters"].get(parts[1])
            if not L:
                self.warnings.append(f"unknown column in {m.group(0)}")
                return m.group(0)
            sheet = f"'{parts[0]}'!"
            which = parts[2] if len(parts) > 2 else "all"
            if which == "last":
                return f"{sheet}${L}${t['last']}"
            if which == "first":
                return f"{sheet}${L}${t['first']}"
            if which == "total" and t.get("totals_row"):
                return f"{sheet}${L}${t['totals_row']}"
            return f"{sheet}${L}${t['first']}:${L}${t['last']}"
        return re.sub(r"\{([^{}]+)\}", rep, f)

    # ------------------------------------------------------------------ title band
    def title_band(self, ws, sh, width_cols):
        title = sh.get("title") or sh["name"]
        ws.row_dimensions[1].height = 30
        c = ws.cell(row=1, column=1, value=title)
        c.font = self.f(size=16, bold=True, color=self.tk["ink"])
        c.alignment = Alignment(vertical="center", horizontal="right" if self.rtl else "left")
        r = 2
        if sh.get("subtitle"):
            s = ws.cell(row=2, column=1, value=sh["subtitle"])
            s.font = self.f(size=9.5, color=self.tk["muted"])
            r = 3
        for col in range(1, max(width_cols, 1) + 1):
            ws.cell(row=r, column=col).border = Border(top=Side(style="medium", color=self.tk["accent"]))
        return r + 1

    # ------------------------------------------------------------------ table sheet
    def sheet_table(self, ws, sh):
        cols = sh.get("columns", [])
        if cols and isinstance(cols[0], str):
            cols = [{"name": c} for c in cols]
        rows = sh.get("rows", [])
        start = self.title_band(ws, sh, len(cols)) + (1 if sh.get("note") else 0)
        if sh.get("note"):
            n = ws.cell(row=start - 1, column=1, value=sh["note"])
            n.font = self.f(size=9, italic=True, color=self.tk["muted"])
        hr = start
        for j, col in enumerate(cols, start=1):
            c = ws.cell(row=hr, column=j, value=col["name"])
            c.font = self.f(bold=True, color=self.tk["on_accent"], size=10)
            c.fill = self.fill(self.tk["accent"])
            c.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center")
        ws.row_dimensions[hr].height = 26
        letters = {col["name"]: get_column_letter(j) for j, col in enumerate(cols, start=1)}
        n = max(len(rows), int(sh.get("blank_rows", 0)))
        first, last = hr + 1, hr + max(n, 1)
        self._first, self._last = first, last
        for i in range(n):
            r = first + i
            row = rows[i] if i < len(rows) else [None] * len(cols)
            for j, col in enumerate(cols, start=1):
                v = row[j - 1] if j - 1 < len(row) else None
                if col.get("formula"):
                    v = self._formula(col["formula"], r, letters)
                elif isinstance(v, str) and col.get("type") in ("date", "datetime") and re.match(r"\d{4}-\d{2}-\d{2}", v):
                    v = dt.datetime.fromisoformat(v)
                elif isinstance(v, str) and v.startswith("="):
                    v = self._formula(v, r, letters)
                c = ws.cell(row=r, column=j, value=v)
                c.number_format = self.fmt(col)
                c.font = self.f(bold=bool(col.get("bold")))
                c.alignment = Alignment(vertical="top", wrap_text=col.get("wrap", col.get("type", "text") == "text" and (col.get("width", 0) or 0) > 30),
                                        horizontal=col.get("align") or ("right" if col.get("type") in ("money", "number", "int", "decimal", "percent", "pct", "score") else None))
                if col.get("input"):
                    c.fill = self.fill("FFF8E1")
        # Excel table object
        ref = f"A{hr}:{get_column_letter(len(cols))}{last}"
        tname = re.sub(r"\W", "", sh.get("table_name") or sh["name"]) or "Table"
        if not re.match(r"[A-Za-z_]", tname):
            tname = "T_" + tname
        tname = tname[:250]
        if tname in [t for s in self.wb.worksheets for t in s.tables]:
            tname += str(len(self.tables))
        tab = Table(displayName=tname, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name=sh.get("table_style", "TableStyleMedium2" if False else "TableStyleLight1"),
                                            showRowStripes=True, showColumnStripes=False)
        ws.add_table(tab)
        # our own banding (works everywhere, also in LibreOffice/Google Sheets)
        band = self.tk["paper_alt"] if self.tk["paper_alt"] != "FFFFFF" else mix(self.tk["accent"], "FFFFFF", 0.94)
        for r in range(first, last + 1):
            for j in range(1, len(cols) + 1):
                c = ws.cell(row=r, column=j)
                c.border = self.thin
                if (r - first) % 2 == 1 and not cols[j - 1].get("input"):
                    c.fill = self.fill(band)
        self.tables[sh["name"]] = {"letters": letters, "first": first, "last": last, "header": hr, "table": tname}
        # totals
        tot = sh.get("totals")
        if tot:
            tr = last + 1
            lab_col = 1
            ws.cell(row=tr, column=lab_col, value=sh.get("totals_label", "الإجمالي" if self.rtl else "Total")).font = self.f(bold=True)
            for name, fn in tot.items():
                L = letters.get(name)
                if not L:
                    self.warnings.append(f"{sh['name']}: totals column '{name}' not found")
                    continue
                code = {"sum": 109, "average": 101, "avg": 101, "count": 103, "max": 104, "min": 105}.get(fn, 109)
                c = ws.cell(row=tr, column=cols.index(next(cc for cc in cols if cc["name"] == name)) + 1,
                            value=f"=SUBTOTAL({code},{L}{first}:{L}{last})")
                c.font = self.f(bold=True)
                c.number_format = self.fmt(next(cc for cc in cols if cc["name"] == name))
            for j in range(1, len(cols) + 1):
                ws.cell(row=tr, column=j).border = Border(top=Side(style="medium", color=self.tk["ink"]))
            self.tables[sh["name"]]["totals_row"] = tr
        # widths
        for j, col in enumerate(cols, start=1):
            w = col.get("width")
            if not w:
                vals = [str(col["name"])] + [str(r[j - 1]) for r in rows[:200] if j - 1 < len(r) and r[j - 1] is not None and not str(r[j - 1]).startswith("=")]
                w = min(48, max(9, max(len(v) for v in vals) * (1.25 if self.rtl else 1.1) + 2))
                if col.get("type") == "money" or (col.get("formula") and col.get("type") in ("number", "decimal")):
                    w = max(w, 17)
                elif col.get("type") in ("number", "int", "decimal", "score") and w < 11:
                    w = 11
            ws.column_dimensions[get_column_letter(j)].width = w
        if sh.get("freeze", True):
            ws.freeze_panes = ws.cell(row=hr + 1, column=1 + int(sh.get("freeze_cols", 0)))
        # validation
        for name, rule in (sh.get("validation") or {}).items():
            L = letters.get(name)
            if not L:
                continue
            rng = f"{L}{first}:{L}{last + int(sh.get('extra_rows', 200))}"
            if isinstance(rule, list):
                dv = DataValidation(type="list", formula1='"' + ",".join(map(str, rule)) + '"', allow_blank=True)
            elif isinstance(rule, dict) and rule.get("list_range"):
                dv = DataValidation(type="list", formula1=rule["list_range"], allow_blank=True)
            else:
                dv = DataValidation(type=rule.get("type", "decimal"), operator=rule.get("operator", "between"),
                                    formula1=str(rule.get("min", 0)), formula2=str(rule.get("max", 100)), allow_blank=True)
                dv.error = rule.get("error", "Value out of range")
                dv.showErrorMessage = True
            ws.add_data_validation(dv)
            dv.add(rng)
        # conditional formatting
        for cf in sh.get("conditional", []):
            L = letters.get(cf.get("column"))
            if not L:
                continue
            rng = f"{L}{first}:{L}{last}"
            k = cf.get("type", "data_bar")
            good, bad, mid = "C6EFCE", "FFC7CE", "FFEB9C"
            if k == "data_bar":
                ws.conditional_formatting.add(rng, DataBarRule(start_type="min", end_type="max", color=self.tk["accent"], showValue=True))
            elif k == "color_scale":
                ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile",
                                                                  mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B")
                                              if not cf.get("reverse") else ColorScaleRule(start_type="min", start_color="63BE7B", mid_type="percentile",
                                                                                           mid_value=50, mid_color="FFEB84", end_type="max", end_color="F8696B"))
            elif k == "icon_set":
                ws.conditional_formatting.add(rng, IconSetRule(cf.get("icons", "3TrafficLights1"), "percent", [0, 33, 67]))
            elif k == "rule":
                fillc = {"good": good, "bad": bad, "warn": mid}.get(cf.get("style", "bad"), cf.get("fill", bad))
                fontc = {"good": "006100", "bad": "9C0006", "warn": "9C5700"}.get(cf.get("style", "bad"), "000000")
                if cf.get("formula"):
                    ws.conditional_formatting.add(f"A{first}:{get_column_letter(len(cols))}{last}" if cf.get("whole_row") else rng,
                                                  FormulaRule(formula=[(lambda x: x[1:] if x.startswith("=") else x)(re.sub(r"(?<![$A-Z])([A-Z]{1,3})(\d+)", r"$\1\2", self._formula(cf["formula"], first, letters))
                                                           if cf.get("whole_row") else self._formula(cf["formula"], first, letters))],
                                                  fill=self.fill(fillc), font=Font(color=fontc)))
                else:
                    ws.conditional_formatting.add(rng, CellIsRule(operator=cf.get("operator", "lessThan"), formula=[str(cf.get("value", 0))],
                                                                  fill=self.fill(fillc), font=Font(color=fontc)))
        for nm, cellref in (sh.get("names") or {}).items():
            self.define_name(nm, ws.title, cellref)

    def _formula(self, f, row, letters, first=None, last=None):
        """Templates: {row} this row · {prev} previous row · {first}/{last} first/last data row ·
        [Column] that column in this row · [Column:all] the whole data range of that column ($-anchored)."""
        first = first or self._first
        last = last or self._last
        f = f.replace("{row}", str(row)).replace("{prev}", str(row - 1)).replace("{first}", str(first)).replace("{last}", str(last))

        def col(m):
            name, _, scope = m.group(1).partition(":")
            if name not in letters:
                return m.group(0)
            L = letters[name]
            return f"${L}${first}:${L}${last}" if scope == "all" else f"{L}{row}"
        f = re.sub(r"\[([^\]]+)\]", col, f)
        return f if f.startswith("=") else "=" + f

    # ------------------------------------------------------------------ inputs (assumptions)
    def sheet_inputs(self, ws, sh):
        r = self.title_band(ws, sh, 3)
        for j, h in enumerate([("البند" if self.rtl else "Input"), ("القيمة" if self.rtl else "Value"), ("ملاحظة" if self.rtl else "Note")], start=1):
            c = ws.cell(row=r, column=j, value=h)
            c.font = self.f(bold=True, color=self.tk["on_accent"])
            c.fill = self.fill(self.tk["accent"])
        for i, fld in enumerate(sh.get("fields", []), start=r + 1):
            ws.cell(row=i, column=1, value=fld["label"]).font = self.f()
            v = ws.cell(row=i, column=2, value=fld.get("value"))
            v.number_format = self.fmt(fld)
            v.fill = self.fill("FFF8E1")
            v.font = self.f(color="1F3FBF", bold=True)          # blue = input (finance convention)
            v.border = Border(bottom=Side(style="thin", color="E0C46C"))
            ws.cell(row=i, column=3, value=fld.get("note", "")).font = self.f(size=9, color=self.tk["muted"])
            if fld.get("name"):
                self.define_name(fld["name"], ws.title, f"$B${i}")
            if fld.get("options"):
                dv = DataValidation(type="list", formula1='"' + ",".join(map(str, fld["options"])) + '"')
                ws.add_data_validation(dv)
                dv.add(f"B{i}")
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 46
        ws.cell(row=r + len(sh.get("fields", [])) + 2, column=1,
                value=("الخلايا الصفراء مدخلات — غيّرها وكل الحسابات هتتحدث" if self.rtl else "Yellow cells are inputs — change them and every calculation updates.")
                ).font = self.f(size=9, italic=True, color=self.tk["muted"])

    def define_name(self, name, sheet, ref):
        absref = ref if "$" in ref else re.sub(r"([A-Z]+)(\d+)", r"$\1$\2", ref)
        q = f"'{sheet}'!{absref}"
        dn = DefinedName(name, attr_text=q)
        try:
            self.wb.defined_names[name] = dn
        except TypeError:  # openpyxl < 3.1
            self.wb.defined_names.append(dn)
        self.names[name] = q

    # ------------------------------------------------------------------ calc block (label / formula rows)
    def sheet_calc(self, ws, sh):
        """A statement-style sheet: rows of label + formula per column (P&L, cash flow, budget summary).
        Formula tokens: {col} this column letter · {prev} previous column letter · {id} line `id` in this
        column · {id@prev} line `id` in the previous column · {id#} the row number of line `id`."""
        cols = sh.get("columns", ["Value"])
        r = self.title_band(ws, sh, len(cols) + 1)
        for j, h in enumerate([""] + cols, start=1):
            c = ws.cell(row=r, column=j, value=h)
            c.font = self.f(bold=True, color=self.tk["on_accent"])
            c.fill = self.fill(self.tk["accent"])
            c.alignment = Alignment(horizontal="center")
        rowmap = {}
        for i, line in enumerate(sh.get("lines", []), start=r + 1):
            if line.get("id"):
                rowmap[line["id"]] = i
        for i, line in enumerate(sh.get("lines", []), start=r + 1):
            lab = ws.cell(row=i, column=1, value=line.get("label", ""))
            strong = line.get("style") in ("total", "header")
            lab.font = self.f(bold=strong, color=self.tk["accent"] if line.get("style") == "header" else self.tk["ink"])
            if line.get("indent"):
                lab.alignment = Alignment(indent=int(line["indent"]))
            vals = line.get("values") or [line.get("formula")] * len(cols) if line.get("formula") or line.get("values") else []
            for j in range(len(cols)):
                v = vals[j] if j < len(vals) else None
                if isinstance(v, str) and ("{col}" in v or "{" in v or v.startswith("=")):
                    L = get_column_letter(j + 2)
                    prevL = get_column_letter(j + 1)
                    v = v.replace("{col}", L).replace("{prev}", prevL)
                    def _ref(m, L=L, prevL=prevL):
                        lid, prev, rowonly = m.group(1), m.group(2), m.group(3)
                        if lid not in rowmap:
                            return m.group(0)
                        if rowonly:
                            return str(rowmap[lid])
                        return f"{prevL if prev else L}{rowmap[lid]}"
                    v = re.sub(r"\{(\w+)(@prev)?(#)?\}", _ref, v)
                    v = v if v.startswith("=") else "=" + v
                c = ws.cell(row=i, column=j + 2, value=v)
                c.number_format = self.fmt(line) if line.get("type") or line.get("format") else self.fmt({"type": sh.get("type", "money")})
                c.font = self.f(bold=strong, color="1F3FBF" if line.get("input") else self.tk["ink"])
                if line.get("input"):
                    c.fill = self.fill("FFF8E1")
                if line.get("style") == "total":
                    c.border = Border(top=Side(style="thin", color=self.tk["ink"]), bottom=Side(style="double", color=self.tk["ink"]))
            if line.get("name"):
                self.define_name(line["name"], ws.title, f"$B${i}")
        ws.column_dimensions["A"].width = max(28, max((len(str(l.get("label", ""))) for l in sh.get("lines", [])), default=10) + 4)
        for j in range(len(cols)):
            ws.column_dimensions[get_column_letter(j + 2)].width = 15
        ws.freeze_panes = ws.cell(row=r + 1, column=2)

    # ------------------------------------------------------------------ dashboard
    def sheet_dashboard(self, ws, sh):
        ws.sheet_view.showGridLines = False
        for col in range(1, 14):
            ws.column_dimensions[get_column_letter(col)].width = 12.5
        r = self.title_band(ws, sh, 12)
        kpis = sh.get("kpis", [])
        span = 3 if len(kpis) <= 4 else 2
        for i, k in enumerate(kpis):
            c0 = 1 + i * span if len(kpis) <= 4 else 1 + (i % 6) * span
            rr = r + 1 + (0 if len(kpis) <= 6 else (i // 6) * 4)
            ws.merge_cells(start_row=rr, start_column=c0, end_row=rr, end_column=c0 + span - 1)
            ws.merge_cells(start_row=rr + 1, start_column=c0, end_row=rr + 1, end_column=c0 + span - 1)
            ws.merge_cells(start_row=rr + 2, start_column=c0, end_row=rr + 2, end_column=c0 + span - 1)
            lab = ws.cell(row=rr, column=c0, value=k.get("label", ""))
            lab.font = self.f(size=9, bold=True, color=self.tk["muted"])
            val = ws.cell(row=rr + 1, column=c0, value=k.get("formula") or k.get("value"))
            val.font = self.f(size=22, bold=True, color=self.tk["accent"])
            val.number_format = self.fmt(k)
            sub = ws.cell(row=rr + 2, column=c0, value=k.get("note", ""))
            sub.font = self.f(size=8.5, color=self.tk["muted"])
            for rrr in range(rr, rr + 3):
                for cc in range(c0, c0 + span):
                    cell = ws.cell(row=rrr, column=cc)
                    cell.fill = self.fill(self.tk["paper_alt"] if self.tk["paper_alt"] != "FFFFFF" else mix(self.tk["accent"], "FFFFFF", 0.93))
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[rr + 1].height = 34
            for cc in range(c0, c0 + span):
                ws.cell(row=rr, column=cc).border = Border(top=Side(style="thick", color=self.tk["accent"]))
        sh.setdefault("_chart_row", r + 1 + 4 * max(1, (len(kpis) + 5) // 6 if len(kpis) > 6 else 1) + 1)

    # ------------------------------------------------------------------ text / notes
    def sheet_text(self, ws, sh):
        r = self.title_band(ws, sh, 8)
        ws.column_dimensions["A"].width = 110
        for line in sh.get("lines", []):
            c = ws.cell(row=r, column=1, value=line if isinstance(line, str) else line.get("text", ""))
            st = {} if isinstance(line, str) else line
            c.font = self.f(bold=st.get("bold", False), size=st.get("size", 10.5), color=self.tk["accent"] if st.get("heading") else self.tk["ink"])
            c.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1

    def contents_sheet(self, sheets):
        ws = self.wb.create_sheet("المحتويات" if self.rtl else "Contents")
        ws.sheet_view.showGridLines = False
        ws.sheet_view.rightToLeft = self.rtl
        ws.column_dimensions["A"].width = 4
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["C"].width = 70
        t = ws.cell(row=2, column=2, value=self.spec.get("title", "Workbook"))
        t.font = self.f(size=20, bold=True)
        if self.spec.get("subtitle"):
            ws.cell(row=3, column=2, value=self.spec["subtitle"]).font = self.f(size=10, color=self.tk["muted"])
        for i, sh in enumerate(sheets, start=5):
            c = ws.cell(row=i, column=2, value=sh.get("title") or sh["name"])
            c.hyperlink = f"#'{sh['name'][:31]}'!A1"
            c.font = self.f(color=self.tk["accent"], bold=True, underline="single")
            ws.cell(row=i, column=3, value=sh.get("description", "")).font = self.f(size=9.5, color=self.tk["ink_soft"])
            ws.row_dimensions[i].height = 20
        self.print_setup(ws, {"name": ws.title, "orientation": "portrait"})

    # ------------------------------------------------------------------ charts
    def chart(self, ws, ch):
        src = self.tables.get(ch.get("data_sheet") or ws.title)
        data_ws = self.wb[ch.get("data_sheet") or ws.title]
        kind = ch.get("type", "bar")
        if src:
            L = src["letters"]
            first, last = src["first"], src["last"]
            cat_col = L.get(ch.get("categories")) if ch.get("categories") in L else ch.get("categories", "A")
            val_cols = [L.get(v, v) for v in ch.get("values", [])]
            hdr = src["header"]
        else:
            first, last, hdr = ch.get("first_row", 2), ch.get("last_row", 10), ch.get("header_row", 1)
            cat_col, val_cols = ch.get("categories", "A"), ch.get("values", ["B"])
        from openpyxl.utils import column_index_from_string as ci
        if kind in ("pie", "donut"):
            c = DoughnutChart(holeSize=55) if kind == "donut" else PieChart()
        elif kind == "line":
            c = LineChart()
        elif kind == "area":
            c = AreaChart()
        elif kind == "scatter":
            c = ScatterChart()
        else:
            c = BarChart()
            c.type = "bar" if kind == "hbar" else "col"
            if kind in ("stacked", "stacked100"):
                c.grouping = "stacked" if kind == "stacked" else "percentStacked"
                c.overlap = 100
            c.gapWidth = 60
        c.title = ch.get("title")
        c.style = 10
        c.height = ch.get("height_cm", 7.5)
        c.width = ch.get("width_cm", 16)
        cats = Reference(data_ws, min_col=ci(cat_col), min_row=first, max_row=last)
        pal = [self.tk["accent"], self.tk["accent_2"], mix(self.tk["accent"], self.tk["ink"], 0.4), mix(self.tk["accent"], "FFFFFF", 0.45),
               mix(self.tk["accent_2"], "FFFFFF", 0.4), self.tk["muted"]]
        for k, vc in enumerate(val_cols):
            ref = Reference(data_ws, min_col=ci(vc), min_row=hdr, max_row=last)
            if kind == "scatter":
                s = Series(Reference(data_ws, min_col=ci(vc), min_row=first, max_row=last), cats, title_from_data=False,
                           title=str(data_ws.cell(row=hdr, column=ci(vc)).value))
                s.marker.symbol = "circle"
                s.graphicalProperties.line.noFill = True
                c.series.append(s)
            else:
                c.add_data(ref, titles_from_data=True)
        if kind != "scatter":
            c.set_categories(cats)
        for k, s in enumerate(c.series):
            col = pal[k % len(pal)]
            if kind in ("line",):
                s.graphicalProperties.line.solidFill = col
                s.graphicalProperties.line.width = 28000
                s.smooth = False
            elif kind not in ("pie", "donut"):
                s.graphicalProperties.solidFill = col
                s.graphicalProperties.line.solidFill = col
        if kind in ("pie", "donut") and c.series:
            from openpyxl.chart.series import DataPoint
            n = last - first + 1
            for i in range(n):
                pt = DataPoint(idx=i)
                pt.graphicalProperties.solidFill = pal[i % len(pal)]
                c.series[0].dPt.append(pt)
            c.dataLabels = DataLabelList()
            c.dataLabels.showPercent = True
            c.dataLabels.showVal = False
            c.dataLabels.showCatName = False
            c.dataLabels.showSerName = False
            c.dataLabels.showLegendKey = False
        elif ch.get("labels"):
            c.dataLabels = DataLabelList()
            c.dataLabels.showVal = True
            c.dataLabels.showCatName = False
            c.dataLabels.showSerName = False
            c.dataLabels.showLegendKey = False
        if len(val_cols) == 1 and kind not in ("pie", "donut"):
            c.legend = None
        anchor = ch.get("anchor") or f"A{ws.max_row + 2}"
        if ws.title in [s.get("name") for s in self.spec.get("sheets", []) if s.get("kind") == "dashboard"] and not ch.get("anchor"):
            sh = next(s for s in self.spec["sheets"] if s["name"] == ws.title)
            i = sh.setdefault("_chart_i", 0)
            row = sh.get("_chart_row", 8) + (i // 2) * 17
            colL = "A" if i % 2 == 0 else "G"
            anchor = f"{colL}{row}"
            sh["_chart_i"] = i + 1
            c.width, c.height = 15, 8
        ws.add_chart(c, anchor)

    # ------------------------------------------------------------------ print
    def print_setup(self, ws, sh):
        ws.page_setup.orientation = sh.get("orientation", "landscape" if ws.max_column > 6 else "portrait")
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.oddFooter.center.text = "&P / &N"
        ws.oddHeader.left.text = self.spec.get("title", "")
        t = self.tables.get(sh["name"])
        if t:
            ws.print_title_rows = f"{t['header']}:{t['header']}"


def recalc_check(xlsx: Path) -> dict:
    """Recalculate with LibreOffice and list formula errors (cell → error)."""
    from pdk.render import soffice
    import shutil
    import subprocess
    import tempfile
    exe = soffice()
    if not exe:
        return {"recalc": "skipped (LibreOffice not installed — open the file in Excel and check for #errors)"}
    tmp = Path(tempfile.mkdtemp(prefix="recalc-"))
    prof = Path(tempfile.mkdtemp(prefix="lo-"))
    subprocess.run([exe, f"-env:UserInstallation={prof.as_uri()}", "--headless", "--calc", "--convert-to", "xlsx",
                    "--outdir", str(tmp), str(xlsx)], capture_output=True, timeout=180)
    out = tmp / xlsx.name
    res = {"recalc": "ok", "errors": []}
    if not out.exists():
        res["recalc"] = "failed to recalculate"
        return res
    import warnings
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(out, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERRORS):
                    res["errors"].append(f"{ws.title}!{c.coordinate} = {c.value}")
    dash = [ws for ws in wb.worksheets if ws.title.lower().startswith(("dash", "لوحة", "ملخص", "summary"))]
    res["sample"] = {f"{ws.title}!{c.coordinate}": c.value for ws in dash[:1] for row in ws.iter_rows(max_row=12) for c in row
                     if isinstance(c.value, (int, float))}
    shutil.rmtree(tmp, ignore_errors=True)
    shutil.rmtree(prof, ignore_errors=True)
    return res


def main(argv=None):
    a = argv or sys.argv[1:]
    args = [x for x in a if not x.startswith("--")]
    if len(args) != 2:
        print(__doc__)
        return 1
    spec_path = Path(args[0])
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    rep = SheetBuilder(spec, spec_path.parent).build(Path(args[1]), recalc="--no-recalc" not in a)
    print("wrote", rep["file"], "| sheets:", ", ".join(rep["sheets"]))
    for w in rep["warnings"]:
        print("WARN:", w)
    if rep.get("errors"):
        print(f"❌ {len(rep['errors'])} formula error(s):")
        for e in rep["errors"][:30]:
            print("  ", e)
        return 2
    print("recalc:", rep.get("recalc"))
    if rep.get("sample"):
        print("dashboard values after recalculation:", json.dumps(rep["sample"], ensure_ascii=False)[:600])
    return 0


if __name__ == "__main__":
    sys.exit(main())
