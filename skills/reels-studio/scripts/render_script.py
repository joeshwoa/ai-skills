#!/usr/bin/env python3
"""Render a reel spec (JSON) into a production script: PDF + DOCX (+ HTML, teleprompter, draft SRT).

    python3 scripts/render_script.py spec.json --out OUT_DIR
        [--formats pdf,docx]            # any of: pdf, docx, html, teleprompter, srt, production (default pdf,docx)
        [--brand neoai]                 # overrides spec["brand"]; slug in brands/ or a .json path
        [--pdf-engine auto|playwright|chrome|soffice]
        [--name FILE_STEM]

Layout = the house style of the user's reel scripts: A4, Cairo, right-aligned for RTL
languages (left for LTR), black headers, editor-facing lines in blue, the talent's spoken
lines in dark red (color_mode "editor_talent"), or every body line in one blue ("mono").

PDF engines, tried in order for "auto": Playwright Chromium -> a Chrome/Chromium/Edge/Brave
binary (CHROME_PATH honoured) -> LibreOffice converting the DOCX. Chromium is preferred:
it shapes Arabic and mixed Arabic/English lines correctly with the bundled font.
The DOCX needs python-docx (bash scripts/setup.sh installs it if missing).
"""
from __future__ import annotations

import argparse
import base64
import html
import os
import re
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from reelspec import (BLOCK_TYPES, FONTS_DIR, PAUSE_SEC, SpecError, as_list,  # noqa: E402
                      base_lang, direction_for, estimate_scene_seconds, has_rtl_chars, line_text,
                      load_brand, load_json, meta_items, speech_seconds)
from pathlib import Path  # noqa: E402


# ----------------------------------------------------------------------------- model
class Para:
    """One rendered paragraph. role: title | meta | note | legend_editor | legend_talent | section |
    section_text | bullet | scene | label | body | spacer.  tone: text | editor | talent"""

    def __init__(self, text: str, role: str, tone: str = "text", weight: int = 400,
                 align: str | None = None, italic: bool = False, space_before: float = 0):
        self.text, self.role, self.tone, self.weight = text, role, tone, weight
        self.align, self.italic, self.space_before = align, italic, space_before


def build_paragraphs(spec: dict, brand: dict) -> list[Para]:
    labels = brand["labels"]
    mono = (spec.get("color_mode") or brand.get("color_mode")) == "mono"
    legend = spec.get("legend", brand.get("legend", True)) and not mono

    def tone_for(audience: str) -> str:
        if audience == "neutral":
            return "text"
        return "mono" if mono else audience

    out: list[Para] = []
    if not spec.get("title"):
        raise SpecError("spec needs a title")
    out.append(Para(spec["title"], "title", weight=500))
    for item in meta_items(spec.get("meta")):
        if isinstance(item, (list, tuple)) and len(item) == 2:
            out.append(Para(f"{item[0]}: {item[1]}", "meta"))
        else:
            out.append(Para(str(item), "meta"))
    if spec.get("intro_note"):
        out.append(Para(spec["intro_note"], "note"))
    if legend:
        out.append(Para(labels["legend_editor"], "legend_editor", "editor", 500, space_before=26))
        out.append(Para(labels["legend_talent"], "legend_talent", "talent", 500))

    for sec in as_list(spec.get("sections")):
        if not isinstance(sec, dict) or not sec.get("heading"):
            raise SpecError("every entry in sections needs a heading")
        out.append(Para(sec["heading"], "section", weight=500, space_before=22))
        tone = tone_for(sec.get("audience", "neutral"))
        for p in as_list(sec.get("paragraphs")):
            out.append(Para(p, "section_text", tone))
        for b in as_list(sec.get("bullets")):
            out.append(Para(b, "bullet", tone))

    if spec.get("sections") and spec.get("scenes"):
        out.append(Para(spec.get("script_heading") or labels["script"], "section", weight=500, space_before=22))

    for i, scene in enumerate(as_list(spec.get("scenes"))):
        if not isinstance(scene, dict):
            raise SpecError(f"scene {i + 1} must be an object with title and blocks")
        title = scene.get("title") or f"SCENE {i + 1:02d}"
        out.append(Para(title, "scene", weight=500, space_before=40 if i or legend or spec.get("sections") else 26))
        if scene.get("duration"):
            out.append(Para(f"Duration: {scene['duration']}", "body"))
        for block in as_list(scene.get("blocks")):
            if not isinstance(block, dict):
                raise SpecError(f"scene '{title}': every block must be an object with a type")
            btype = block.get("type")
            if btype not in BLOCK_TYPES:
                raise SpecError(f"scene '{title}': unknown block type '{btype}' (allowed: {', '.join(BLOCK_TYPES)})")
            default_label, audience = BLOCK_TYPES[btype]
            if btype in ("dialogue", "voiceover"):
                label = block.get("label") or block.get("speaker") or default_label
                # In the house style the speaker label is black; the words are the talent's colour.
                out.append(Para(f"{label}:", "label", "text" if not mono else "text", 500, space_before=6))
            else:
                label = block.get("label") or default_label
                out.append(Para(f"{label}:", "label", tone_for(audience), 500, space_before=6))
            for item in as_list(block.get("lines")):
                kind, text = line_text(item)
                out.append(Para(text, "body", tone_for(audience), italic=(kind == "direction")))

    ef = spec.get("end_frame")
    if ef is None and brand.get("end_frame_default"):
        ef = {"lines": brand["end_frame_default"]}
    if isinstance(ef, (list, str)):
        ef = {"lines": as_list(ef)}
    if ef:
        out.append(Para(labels["end_frame"], "scene", weight=500, space_before=40))
        for line in as_list(ef.get("lines")):
            out.append(Para(line, "body", "text", 500))
        for line in as_list(ef.get("notes")):
            out.append(Para(line, "body", tone_for("editor")))

    if spec.get("editor_notes"):
        out.append(Para(labels["editor_notes"], "scene", tone_for("editor"), 500, space_before=40))
        for line in as_list(spec["editor_notes"]):
            out.append(Para(line, "body", tone_for("editor")))

    posting = spec.get("posting")
    if isinstance(posting, dict) and posting and spec.get("render_posting_pack", True):
        out.append(Para(labels["posting"], "scene", weight=500, space_before=40))
        for key, lab in (("cover_text", "Cover text"), ("title", "Title (YouTube Shorts)"), ("caption", "Caption"), ("first_comment", "First comment"),
                         ("best_time", "Best time to post")):
            if posting.get(key):
                out.append(Para(f"{lab}:", "label", "text", 500, space_before=6))
                for ln in str(posting[key]).split("\n"):
                    out.append(Para(ln, "body", tone_for("editor")))
        if posting.get("hashtags"):
            out.append(Para("Hashtags:", "label", "text", 500, space_before=6))
            out.append(Para(" ".join(map(str, as_list(posting["hashtags"]))), "body", tone_for("editor")))
        if posting.get("alt_hooks"):
            out.append(Para("A/B hooks:", "label", "text", 500, space_before=6))
            for h in as_list(posting["alt_hooks"]):
                out.append(Para(h, "body", tone_for("talent")))
    return out


# ----------------------------------------------------------------------------- HTML
def font_faces(brand: dict) -> str:
    fam = brand["font"]["family"]
    weights = {"Cairo-Regular.ttf": 400, "Cairo-Medium.ttf": 500, "Cairo-SemiBold.ttf": 600, "Cairo-Bold.ttf": 700}
    css = []
    for fname in brand["font"].get("files") or []:
        path = fname if os.path.isabs(fname) else os.path.join(FONTS_DIR, fname)
        if not os.path.exists(path):
            continue
        data = base64.b64encode(open(path, "rb").read()).decode()
        w = weights.get(os.path.basename(fname), 400)
        css.append(f"@font-face{{font-family:'{fam}';font-weight:{w};font-style:normal;"
                   f"src:url(data:font/ttf;base64,{data}) format('truetype');}}")
    return "\n".join(css)


def to_html(spec: dict, brand: dict, paras: list[Para]) -> str:
    d = direction_for(spec)
    lang = spec.get("language") or "en"
    c, s = brand["colors"], brand["sizes_pt"]
    main_align = "right" if d == "rtl" else "left"
    fam = brand["font"]["family"]
    stack = ", ".join([f"'{fam}'"] + [f"'{f}'" if " " in f else f for f in brand["font"]["fallbacks"]])
    size_for = {"title": s["title"], "meta": s["meta"], "note": s["body"], "legend_editor": s["legend"],
                "legend_talent": s["legend"], "section": s["section"], "section_text": s["body"],
                "bullet": s["body"], "scene": s["scene"], "label": s["label"], "body": s["body"]}
    color = {"text": c["text"], "editor": c["editor"], "talent": c["talent"], "mono": c["mono"]}
    rows = []
    for p in paras:
        align = p.align or ("left" if p.role.startswith("legend") else main_align)
        style = [f"font-size:{size_for.get(p.role, s['body'])}pt", f"color:#{color[p.tone]}",
                 f"font-weight:{p.weight}", f"text-align:{align}"]
        if p.italic:
            style.append("font-style:italic")
        if p.space_before:
            style.append(f"margin-top:{p.space_before}pt")
        text = html.escape(p.text)
        if p.role == "bullet":
            text = "• " + text
        # A line with any Arabic/Hebrew in it is laid out RTL (so "Camera closer شوية" reads right),
        # a Latin-only line LTR (so "Visual:" keeps its colon on the right side of the word).
        pdir = "rtl" if has_rtl_chars(p.text) else "ltr"
        rows.append(f'<p dir="{pdir}" class="{p.role}" style="{";".join(style)}">{text}</p>')
    margin = brand["page"].get("margin_mm", 12)
    return f"""<!doctype html>
<html lang="{html.escape(lang)}" dir="{d}"><head><meta charset="utf-8">
<title>{html.escape(spec['title'])}</title>
<style>
{font_faces(brand)}
@page {{ size: {brand['page'].get('size', 'A4')}; margin: {margin}mm; }}
html, body {{ margin: 0; padding: 0; background: #fff; }}
body {{ font-family: {stack}; line-height: 1.32; -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
p {{ margin: 0 0 11pt 0; unicode-bidi: isolate; white-space: pre-wrap; break-inside: avoid; }}
p.title {{ margin-bottom: 14pt; line-height: 1.25; }}
p.meta {{ margin-bottom: 9pt; }}
p.scene {{ break-after: avoid; }}
p.label {{ break-after: avoid; }}
@media screen {{ body {{ max-width: 186mm; margin: 12mm auto; }} }}
</style></head>
<body>
{chr(10).join(rows)}
</body></html>
"""


# ----------------------------------------------------------------------------- DOCX
RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
             "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
             "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
             "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]
def to_docx(spec: dict, brand: dict, paras: list[Para], path: str) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor
    except ImportError as e:
        raise SpecError("DOCX output needs python-docx — run: bash scripts/setup.sh  (or pip install python-docx)") from e

    d = direction_for(spec)
    c, s = brand["colors"], brand["sizes_pt"]
    fam, fam_medium = brand["font"]["family"], brand["font"].get("medium") or brand["font"]["family"]
    size_for = {"title": s["title"], "meta": s["meta"], "note": s["body"], "legend_editor": s["legend"],
                "legend_talent": s["legend"], "section": s["section"], "section_text": s["body"],
                "bullet": s["body"], "scene": s["scene"], "label": s["label"], "body": s["body"]}

    doc = Document()
    sec = doc.sections[0]
    if brand["page"].get("size", "A4").upper() == "A4":
        sec.page_width, sec.page_height = Mm(210), Mm(297)
    m = Mm(brand["page"].get("margin_mm", 12))
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = m
    lang_tag = spec.get("language") or "en"

    def set_fonts(run, family: str, size: float, bold: bool, italic: bool, rgb: str):
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor.from_string(rgb)
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), family)
        # complex-script (Arabic, Hebrew…) runs read the *Cs properties, not the Latin ones
        for tag, on in (("w:bCs", bold), ("w:iCs", italic)):
            el = OxmlElement(tag)
            el.set(qn("w:val"), "1" if on else "0")
            rpr.append(el)
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(int(round(size * 2))))
        rpr.append(szcs)
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "en-US")
        lang.set(qn("w:bidi"), lang_tag if base_lang(lang_tag) in ("ar", "fa", "he", "ur") else "ar-EG")
        rpr.append(lang)
        # Word validates rPr children against a fixed sequence; put everything in schema order.
        order = {qn("w:" + t): i for i, t in enumerate(RPR_ORDER)}
        kids = sorted(list(rpr), key=lambda el: order.get(el.tag, len(order)))
        for el in kids:
            rpr.remove(el)
        for el in kids:
            rpr.append(el)

    for p in paras:
        para = doc.add_paragraph()
        fmt = para.paragraph_format
        fmt.space_after = Pt(9 if p.role == "meta" else 11)
        fmt.space_before = Pt(p.space_before)
        fmt.line_spacing = 1.2
        rtl_line = d == "rtl" and has_rtl_chars(p.text)
        if rtl_line:
            ppr = para._p.get_or_add_pPr()
            ppr.insert(0, OxmlElement("w:bidi"))
        legend = p.role.startswith("legend")
        want_right = d == "rtl" and not legend
        # In a bidi paragraph Word and LibreOffice read left/right as start/end, so a right-aligned
        # Arabic line is "left" there; a Latin-only line in an RTL script stays physically "right".
        if rtl_line:
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT if want_right else WD_ALIGN_PARAGRAPH.RIGHT
        else:
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if want_right else WD_ALIGN_PARAGRAPH.LEFT
        text = p.text
        if p.role == "bullet":
            text = "• " + text
        family = fam_medium if p.weight == 500 else fam
        rgb = {"text": c["text"], "editor": c["editor"], "talent": c["talent"], "mono": c["mono"]}[p.tone]
        run = para.add_run(text)
        set_fonts(run, family, size_for.get(p.role, s["body"]), p.weight >= 600, p.italic, rgb)
    doc.core_properties.title = spec["title"]
    # python-docx's default settings.xml has <w:zoom> without the required w:percent
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    doc.save(path)


# ----------------------------------------------------------------------------- PDF engines
CHROME_CANDIDATES = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
    "/Applications/Brave Browser.app/Contents/MacOS/Brave Browser",
    "/opt/pw-browsers/chromium/chrome-linux/chrome",
]
CHROME_NAMES = ["google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "microsoft-edge", "brave-browser"]


def find_chrome() -> str | None:
    env = os.environ.get("CHROME_PATH")
    if env and os.path.exists(env):
        return env
    for p in CHROME_CANDIDATES:
        if os.path.exists(p):
            return p
    for n in CHROME_NAMES:
        w = shutil.which(n)
        if w:
            return w
    return None


def find_soffice() -> str | None:
    for p in ["/Applications/LibreOffice.app/Contents/MacOS/soffice"]:
        if os.path.exists(p):
            return p
    return shutil.which("soffice") or shutil.which("libreoffice")


def pdf_playwright(html_path: str, pdf_path: str) -> None:
    from playwright.sync_api import sync_playwright  # noqa
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        page = browser.new_page()
        page.goto(Path(html_path).resolve().as_uri())
        page.wait_for_load_state("networkidle")
        page.evaluate("document.fonts.ready")
        page.pdf(path=pdf_path, prefer_css_page_size=True, print_background=True)
        browser.close()


def pdf_chrome(html_path: str, pdf_path: str) -> None:
    exe = find_chrome()
    if not exe:
        raise RuntimeError("no Chrome/Chromium/Edge/Brave binary found (set CHROME_PATH)")
    with tempfile.TemporaryDirectory() as prof:
        cmd = [exe, "--headless=new", "--disable-gpu", "--no-sandbox", f"--user-data-dir={prof}",
               "--no-pdf-header-footer", "--virtual-time-budget=4000",
               f"--print-to-pdf={os.path.abspath(pdf_path)}", Path(html_path).resolve().as_uri()]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    if not os.path.exists(pdf_path) or os.path.getsize(pdf_path) == 0:
        raise RuntimeError(f"chrome produced no PDF: {r.stderr[-400:]}")


def pdf_soffice(docx_path: str, pdf_path: str) -> None:
    exe = find_soffice()
    if not exe:
        raise RuntimeError("LibreOffice (soffice) not found")
    outdir = tempfile.mkdtemp()
    env = dict(os.environ)
    if sys.platform != "darwin":
        # Let LibreOffice see the bundled fonts without installing them: a fontconfig file that
        # includes the system config plus this skill's font folder.
        conf = os.path.join(outdir, "fonts.conf")
        with open(conf, "w", encoding="utf-8") as f:
            f.write('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>'
                    f'<include ignore_missing="yes">/etc/fonts/fonts.conf</include><dir>{html.escape(FONTS_DIR)}</dir>'
                    f'<cachedir>{html.escape(os.path.join(outdir, "fc-cache"))}</cachedir></fontconfig>')
        env["FONTCONFIG_FILE"] = conf
    elif not any(os.path.exists(os.path.join(d, "Cairo-Regular.ttf"))
                 for d in (os.path.expanduser("~/Library/Fonts"), "/Library/Fonts")):
        print("note: Cairo isn't installed, so LibreOffice will substitute a font — run "
              "`bash scripts/setup.sh --fonts` (or install Chrome, the preferred PDF engine).", file=sys.stderr)
    profile = os.path.join(outdir, "lo-profile")
    r = subprocess.run([exe, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", "pdf",
                        "--outdir", outdir, docx_path], capture_output=True, text=True, timeout=180, env=env)
    produced = os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
    if not os.path.exists(produced):
        raise RuntimeError(f"soffice produced no PDF: {r.stderr[-400:]}")
    shutil.move(produced, pdf_path)


def make_pdf(engine: str, html_path: str, docx_path: str | None, pdf_path: str) -> str:
    order = ["playwright", "chrome", "soffice"] if engine == "auto" else [engine]
    errors = []
    for e in order:
        if os.path.exists(pdf_path):  # never report a leftover file from an earlier run as success
            os.remove(pdf_path)
        try:
            if e == "playwright":
                pdf_playwright(html_path, pdf_path)
            elif e == "chrome":
                pdf_chrome(html_path, pdf_path)
            elif e == "soffice":
                if not docx_path:
                    raise RuntimeError("soffice engine needs the DOCX (add docx to --formats)")
                pdf_soffice(docx_path, pdf_path)
            if not os.path.exists(pdf_path) or os.path.getsize(pdf_path) < 500:
                raise RuntimeError("engine finished but wrote no PDF")
            return e
        except Exception as ex:  # try the next engine
            errors.append(f"{e}: {str(ex).splitlines()[0] if str(ex) else type(ex).__name__}")
    raise SpecError("no PDF engine worked —\n  " + "\n  ".join(errors) +
                    "\nInstall one: Google Chrome, `pip install playwright && playwright install chromium`, or LibreOffice.")


# ----------------------------------------------------------------------------- teleprompter / SRT
def teleprompter(spec: dict) -> str:
    out = [spec["title"], ""]
    for scene in as_list(spec.get("scenes")):
        lines = []
        for block in as_list(scene.get("blocks")):
            if block.get("type") in ("dialogue", "voiceover"):
                for item in as_list(block.get("lines")):
                    kind, text = line_text(item)
                    lines.append(f"[{text}]" if kind in ("pause", "direction") else text)
        if lines:
            out.append(f"— {scene.get('title', '')}")
            out.extend(lines)
            out.append("")
    return "\n".join(out).strip() + "\n"


def wrap_caption(text: str, max_chars: int, by_char: bool = False) -> list:
    if by_char:  # scripts without spaces (zh/ja/th): cut by characters
        t = text.replace(" ", "")
        return [t[i:i + max_chars] for i in range(0, len(t), max_chars)] or [""]
    words, chunks, cur = text.split(), [], ""
    for w in words:
        if cur and len(cur) + 1 + len(w) > max_chars:
            chunks.append(cur)
            cur = w
        else:
            cur = f"{cur} {w}".strip()
    if cur:
        chunks.append(cur)
    return chunks


def srt(spec: dict) -> str:
    lang = spec.get("language") or "en"
    wps = (spec.get("brief") or {}).get("words_per_sec")
    by_char = base_lang(lang) in ("zh", "ja", "th")
    max_chars = 16 if by_char else (40 if base_lang(lang) == "ar" else 42)
    t, idx, cues = 0.0, 1, []

    def ts(x: float) -> str:
        ms = int(round(x * 1000))
        h, ms = divmod(ms, 3600000)
        m, ms = divmod(ms, 60000)
        s_, ms = divmod(ms, 1000)
        return f"{h:02d}:{m:02d}:{s_:02d},{ms:03d}"

    for scene in as_list(spec.get("scenes")):
        start_scene = t
        for block in as_list(scene.get("blocks")):
            if block.get("type") not in ("dialogue", "voiceover"):
                continue
            for item in as_list(block.get("lines")):
                kind, text = line_text(item)
                if kind == "pause":
                    t += PAUSE_SEC
                    continue
                if kind != "line":
                    continue
                clean = text.strip().strip('"“”').strip()
                for chunk in wrap_caption(clean, max_chars, by_char):
                    dur = max(0.8, speech_seconds(chunk, lang, wps))
                    cues.append(f"{idx}\n{ts(t)} --> {ts(t + dur)}\n{chunk}\n")
                    idx += 1
                    t += dur
        if isinstance(scene.get("duration_sec"), (int, float)):
            t = max(t, start_scene + float(scene["duration_sec"]))
    header = ""  # SRT has no comment syntax; the estimate warning is printed by the CLI instead
    return header + "\n".join(cues)


def production_text(spec: dict) -> str:
    """Plain-text script in the labelled inline format flow-studio-director's script_breakdown.py
    reads (scene headings, `Name: line`, SFX:/Music:/Text:/VO:/Visual:, (parentheticals))."""
    talent = None
    for item in meta_items(spec.get("meta")):
        if isinstance(item, (list, tuple)) and len(item) == 2 and str(item[0]).lower() in ("talent", "cast"):
            talent = str(item[1]).split("/")[0].split(",")[0].strip()
    label_for = {"visual": "Visual", "camera": "Camera", "broll": "Visual", "sfx": "SFX", "music": "Music",
                 "onscreen_text": "Text", "transition": "Note", "note": "Note"}
    out = [f"Title: {spec['title']}"]
    brief = spec.get("brief") or {}
    for k, lab in (("platforms", "Platform"), ("audience", "Target"), ("duration_sec", "Duration"), ("format", "Format")):
        if brief.get(k):
            v = brief[k]
            if k == "duration_sec":
                v = (f"{v[0]}–{v[1]} sec" if isinstance(v, list) else f"{v} sec")
            elif isinstance(v, list):
                v = ", ".join(map(str, v))
            out.append(f"{lab}: {v}")
    out.append(f"Note: language {spec.get('language')}")
    for n in as_list(spec.get("editor_notes")):  # editor notes belong to the whole piece → preamble
        out.append(f"Note: {n}")
    for i, scene in enumerate(as_list(spec.get("scenes")), 1):
        title = scene.get("title") or f"SCENE {i:02d}"
        if not re.match(r"(?i)^\s*(scene|sc\.?|مشهد|المشهد)\b", title):
            title = f"SCENE {i:02d} — {title}"
        out += ["", title]
        if scene.get("duration"):
            out.append(f"Duration: {scene['duration']}")
        for block in as_list(scene.get("blocks")):
            t = block.get("type")
            for item in as_list(block.get("lines")):
                kind, text = line_text(item)
                text = text.strip()
                if t in ("dialogue", "voiceover"):
                    if kind in ("pause", "direction"):
                        out.append(f"({text.rstrip('.')})")
                    elif t == "voiceover":
                        out.append(f"VO: {text.strip(chr(34) + '“”')}")
                    else:
                        who = block.get("speaker") or talent or "Talent"
                        out.append(f"{who}: {text.strip(chr(34) + '“”')}")
                elif t == "direction":
                    out.append(f"({text})")
                else:
                    out.append(f"{label_for.get(t, 'Note')}: {text}")
    ef = spec.get("end_frame")
    if isinstance(ef, (list, str)):
        ef = {"lines": as_list(ef)}
    if ef:
        # A real scene heading, so the reader doesn't take "END FRAME" for a character cue.
        n = len(as_list(spec.get("scenes"))) + 1
        out += ["", f"SCENE {n:02d} — END FRAME"] + [f"Text: {l}" for l in as_list(ef.get("lines"))] + \
               [f"Note: {l}" for l in as_list(ef.get("notes"))]
    return "\n".join(out) + "\n"


# ----------------------------------------------------------------------------- main
def safe_stem(title: str) -> str:
    stem = re.sub(r"[^\w\-]+", "_", title, flags=re.UNICODE).strip("_")
    return (stem[:80] or "reel_script")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec")
    ap.add_argument("--out", required=True)
    ap.add_argument("--formats", default="pdf,docx")
    ap.add_argument("--brand")
    ap.add_argument("--pdf-engine", default="auto", choices=["auto", "playwright", "chrome", "soffice"])
    ap.add_argument("--name")
    a = ap.parse_args(argv)

    try:
        spec = load_json(a.spec)
        brand = load_brand(a.brand or spec.get("brand"))
        formats = {f.strip() for f in a.formats.split(",") if f.strip()}
        unknown = formats - {"pdf", "docx", "html", "teleprompter", "srt", "production"}
        if unknown:
            raise SpecError(f"unknown format(s): {', '.join(sorted(unknown))}")
        os.makedirs(a.out, exist_ok=True)
        if not isinstance(spec, dict) or not spec.get("title"):
            raise SpecError("spec needs a title")
        stem = a.name or safe_stem(spec["title"])
        paras = build_paragraphs(spec, brand)
        written = []

        html_path = os.path.join(a.out, stem + ".html")
        need_html = "html" in formats or "pdf" in formats
        if need_html:
            with open(html_path, "w", encoding="utf-8") as f:
                f.write(to_html(spec, brand, paras))
            if "html" in formats:
                written.append(html_path)

        docx_path = None
        if "docx" in formats or ("pdf" in formats and a.pdf_engine == "soffice"):
            docx_path = os.path.join(a.out, stem + ".docx")
            to_docx(spec, brand, paras, docx_path)
            if "docx" in formats:
                written.append(docx_path)

        engine = None
        if "pdf" in formats:
            pdf_path = os.path.join(a.out, stem + ".pdf")
            if a.pdf_engine == "auto" and not docx_path:
                try:
                    docx_path = os.path.join(tempfile.mkdtemp(), stem + ".docx")
                    to_docx(spec, brand, paras, docx_path)
                except SpecError:
                    docx_path = None
            engine = make_pdf(a.pdf_engine, html_path, docx_path, pdf_path)
            written.append(pdf_path)
        if need_html and "html" not in formats and os.path.exists(html_path):
            os.remove(html_path)

        if "teleprompter" in formats:
            p = os.path.join(a.out, stem + "_teleprompter.txt")
            open(p, "w", encoding="utf-8").write(teleprompter(spec))
            written.append(p)
        if "production" in formats:
            p = os.path.join(a.out, stem + "_for_production.txt")
            open(p, "w", encoding="utf-8").write(production_text(spec))
            written.append(p)
        if "srt" in formats:
            p = os.path.join(a.out, stem + "_captions_ESTIMATED.srt")
            open(p, "w", encoding="utf-8").write(srt(spec))
            written.append(p)
            print("note: SRT timings are estimates from reading speed — re-time against the real "
                  "recording (creative-studio audio_transcribe / video_captions) before publishing.")

        total = sum(estimate_scene_seconds(sc, spec.get("language") or "en",
                                           (spec.get("brief") or {}).get("words_per_sec"))
                    for sc in as_list(spec.get("scenes")))
        for w in written:
            print("wrote", w)
        if engine:
            print("pdf engine:", engine)
        print(f"estimated runtime: {total:.0f} s")
        return 0
    except SpecError as e:
        print("error:", e, file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
